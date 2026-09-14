from __future__ import annotations

import csv
import json
import queue
import re
import threading
import time
from pathlib import Path

from PySide6.QtCore import QProcess, QTimer
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QFileDialog,
    QFrame,
    QGridLayout,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPlainTextEdit,
    QPushButton,
    QScrollArea,
    QSpinBox,
    QStyle,
    QSystemTrayIcon,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QVBoxLayout,
    QWidget,
)

from .core import (
    APP_DIR,
    APP_NAME,
    MAX_IMPORT_BYTES,
    PROFILES_DIR,
    app_data_dir,
    atomic_write_text,
    classify_error,
    command_with_destination,
    detect_config_path,
    parse_text_database,
    quote_arg_for_preview,
    read_text_safely,
    redact_sensitive_argv,
    redact_sensitive_text,
    safe_expand_path,
    safe_filename,
    safe_int,
    sanitize_spreadsheet_cell,
    split_command,
    timestamp_slug,
)
from .models import DownloadJob, JobResult
from .workers import DownloadWorker


class AdvancedToolsMixin:
    def _legacy_open_command_builder(self) -> None:
        """Retained for migration reference; the public UI uses ComposerMixin."""
        ind = self._ui_is_indonesian()
        dlg = QDialog(self)
        dlg.setWindowTitle("Pembuat Command" if ind else "Command Builder")
        dlg.resize(900, 660)
        dlg.setMinimumSize(780, 560)
        root = QVBoxLayout(dlg)
        root.setContentsMargins(10, 10, 10, 10)
        root.setSpacing(8)

        intro = QLabel(
            "Buat command gallery-dl dengan aman. Builder ini hanya menyusun command, bukan menjalankan download."
            if ind else
            "Build a safe gallery-dl command. This builder only creates commands, it does not start downloads."
        )
        intro.setObjectName("subtle")
        intro.setWordWrap(True)
        root.addWidget(intro)

        tabs = QTabWidget()
        root.addWidget(tabs, 1)

        def bt(en: str, id_text: str) -> str:
            return id_text if ind else en

        # ----- Build tab ----- #
        build_page = QWidget()
        build_l = QGridLayout(build_page)
        build_l.setContentsMargins(8, 8, 8, 8)
        build_l.setHorizontalSpacing(10)
        build_l.setVerticalSpacing(7)
        tabs.addTab(build_page, "Build" if not ind else "Buat")

        url = QLineEdit()
        url.setPlaceholderText("https://...")
        dest = QLineEdit()
        dest.setPlaceholderText(r"D:\Rips\Creator A atau ~/Downloads/gallery-dl" if ind else r"D:\Rips\Creator A or ~/Downloads/gallery-dl")
        exact_dir = QLineEdit()
        exact_dir.setPlaceholderText(bt("Optional exact directory for -D/--directory", "Folder exact opsional untuk -D/--directory"))
        filename = QLineEdit()
        filename.setPlaceholderText(bt("Optional filename format, e.g. {id}_{filename}.{extension}", "Format nama file opsional, contoh {id}_{filename}.{extension}"))
        tag = QLineEdit()
        tag.setPlaceholderText(bt("Creator A / Batch 01 / Project name", "Creator A / Batch 01 / Nama proyek"))
        notes = QLineEdit()
        notes.setPlaceholderText(bt("Optional notes for CSV/XLSX", "Catatan opsional untuk CSV/XLSX"))

        preset = QComboBox()
        preset_items_en = [
            "Basic download",
            "Download to selected folder",
            "Safe slow mode",
            "Preview only / simulate",
            "Get direct media URLs",
            "Range 1-20",
            "Use browser cookies",
            "Archive skip downloaded",
            "Metadata + info JSON",
            "Windows-safe filenames",
            "Keyword discovery (-K)",
            "Extractor debug info (-E)",
            "JSON metadata audit",
            "Rate-limit friendly",
            "Write logs + error file",
            "ZIP archive output",
            "CBZ manga archive",
            "Date filter example",
            "Tag whitelist example",
            "File size filter example",
            "SSL certificate bypass",
            "HypnoHub cautious mode",
            "Rule34 cautious mode",
            "Booru tag batch",
            "Resume large batch",
        ]
        preset_items_id = [
            "Download dasar",
            "Download ke folder terpilih",
            "Mode lambat aman",
            "Pratinjau saja / simulasi",
            "Ambil URL media langsung",
            "Rentang 1-20",
            "Pakai cookies browser",
            "Arsip skip yang sudah diunduh",
            "Metadata + info JSON",
            "Nama file aman Windows",
            "Cari keyword (-K)",
            "Info debug extractor (-E)",
            "Audit metadata JSON",
            "Ramah rate-limit",
            "Tulis log + file error",
            "Output arsip ZIP",
            "Arsip manga CBZ",
            "Contoh filter tanggal",
            "Contoh whitelist tag",
            "Contoh filter ukuran file",
            "Bypass sertifikat SSL",
            "Mode hati-hati HypnoHub",
            "Mode hati-hati Rule34",
            "Batch tag booru",
            "Lanjutkan batch besar",
        ]
        preset.addItems(preset_items_id if ind else preset_items_en)
        preset.setFixedHeight(28)

        def preset_key(name: str) -> str:
            return dict(zip(preset_items_id, preset_items_en, strict=False)).get(name, name)

        def add_label(row: int, text: str) -> None:
            label = QLabel(text)
            label.setObjectName("fieldLabel")
            build_l.addWidget(label, row, 0)

        add_label(0, bt("Preset", "Preset"))
        build_l.addWidget(preset, 0, 1, 1, 3)
        add_label(1, bt("URL", "URL"))
        build_l.addWidget(url, 1, 1, 1, 3)
        add_label(2, bt("Destination (-d)", "Tujuan (-d)"))
        build_l.addWidget(dest, 2, 1, 1, 3)
        add_label(3, bt("Exact directory (-D)", "Folder exact (-D)"))
        build_l.addWidget(exact_dir, 3, 1, 1, 3)
        add_label(4, bt("Filename (-f)", "Nama file (-f)"))
        build_l.addWidget(filename, 4, 1, 1, 3)
        add_label(5, bt("Tag", "Tag"))
        build_l.addWidget(tag, 5, 1, 1, 3)
        add_label(6, bt("Notes", "Catatan"))
        build_l.addWidget(notes, 6, 1, 1, 3)
        build_l.setColumnStretch(1, 1)
        build_l.setColumnStretch(3, 1)
        build_l.setRowStretch(7, 1)

        # ----- Options tab ----- #
        opt_scroll = QScrollArea()
        opt_scroll.setWidgetResizable(True)
        opt_scroll.setFrameShape(QFrame.NoFrame)
        opt_page = QWidget()
        opt_l = QGridLayout(opt_page)
        opt_l.setContentsMargins(8, 8, 8, 8)
        opt_l.setHorizontalSpacing(10)
        opt_l.setVerticalSpacing(7)
        opt_scroll.setWidget(opt_page)
        tabs.addTab(opt_scroll, "Options" if not ind else "Opsi")

        range_edit = QLineEdit()
        range_edit.setPlaceholderText("1-20 atau 1:50:2" if ind else "1-20 or 1:50:2")
        post_range_edit = QLineEdit()
        post_range_edit.setPlaceholderText("1-10")
        child_range_edit = QLineEdit()
        child_range_edit.setPlaceholderText("1-5")
        date_after = QLineEdit()
        date_after.setPlaceholderText("2026-01-01")
        date_before = QLineEdit()
        date_before.setPlaceholderText("2026-12-31")
        sleep_edit = QLineEdit()
        sleep_edit.setPlaceholderText("2.0-5.0")
        sleep_request = QLineEdit()
        sleep_request.setPlaceholderText("1.0")
        sleep_429 = QLineEdit()
        sleep_429.setPlaceholderText("60-300")
        sleep_skip = QLineEdit()
        sleep_skip.setPlaceholderText("0.5-2.0")
        sleep_extractor = QLineEdit()
        sleep_extractor.setPlaceholderText("1.0-3.0")
        sleep_retries = QLineEdit()
        sleep_retries.setPlaceholderText("60-300")
        retries_spin = QSpinBox()
        retries_spin.setRange(0, 99)
        retries_spin.setValue(0)
        retries_spin.setToolTip("0 = jangan tambahkan -R/--retries" if ind else "0 = do not add -R/--retries")
        limit_rate = QLineEdit()
        limit_rate.setPlaceholderText("500k, 2M, 800k-2M")
        chunk_size = QLineEdit()
        chunk_size.setPlaceholderText("256k, 1M")
        filesize_min = QLineEdit()
        filesize_min.setPlaceholderText("500k, 2M")
        filesize_max = QLineEdit()
        filesize_max.setPlaceholderText("10M, 250M")
        timeout_spin = QSpinBox()
        timeout_spin.setRange(0, 999)
        timeout_spin.setValue(0)
        timeout_spin.setSuffix(" s")
        archive_path = QLineEdit()
        archive_path.setPlaceholderText(str(APP_DIR / "archive.sqlite3"))
        abort_after = QLineEdit()
        abort_after.setPlaceholderText("3 atau 3:tag" if ind else "3 or 3:tag")
        terminate_after = QLineEdit()
        terminate_after.setPlaceholderText("5")
        cookies_browser = QComboBox()
        cookies_items_en = ["none", "use GUI selection", "chrome", "firefox", "edge", "brave", "chromium", "opera"]
        cookies_items_id = ["tidak ada", "pakai pilihan GUI", "chrome", "firefox", "edge", "brave", "chromium", "opera"]
        cookies_browser.addItems(cookies_items_id if ind else cookies_items_en)
        cookies_file = QLineEdit()
        cookies_file.setPlaceholderText("Path cookies.txt opsional untuk -C/--cookies" if ind else "Optional cookies.txt path for -C/--cookies")
        cookies_export = QLineEdit()
        cookies_export.setPlaceholderText("export-cookies.txt")
        proxy = QLineEdit()
        proxy.setPlaceholderText("http://127.0.0.1:8080 atau socks5://..." if ind else "http://127.0.0.1:8080 or socks5://...")
        user_agent = QLineEdit()
        user_agent.setPlaceholderText("User-Agent opsional" if ind else "Optional User-Agent")
        option_kv = QLineEdit()
        option_kv.setPlaceholderText('key=value; opsi_lain=value, contoh extractor.sleep=2; browser=firefox' if ind else 'key=value; other=value, e.g. extractor.sleep=2; browser=firefox')
        filter_expr = QLineEdit()
        filter_expr.setPlaceholderText('image_width >= 1000 and extension in ("jpg", "png")')
        post_filter_expr = QLineEdit()
        post_filter_expr.setPlaceholderText('date >= "2026-01-01"')
        child_filter_expr = QLineEdit()
        child_filter_expr.setPlaceholderText('subcategory != "comments"')
        blacklist = QLineEdit()
        blacklist.setPlaceholderText('pixiv,user:* atau daftar kategori' if ind else 'pixiv,user:* or category list')
        whitelist = QLineEdit()
        whitelist.setPlaceholderText('pixiv,danbooru atau daftar kategori' if ind else 'pixiv,danbooru or category list')
        tags_blacklist = QLineEdit()
        tags_blacklist.setPlaceholderText('tag1,tag2 atau path/ke/list.txt' if ind else 'tag1,tag2 or path/to/list.txt')
        tags_whitelist = QLineEdit()
        tags_whitelist.setPlaceholderText('tag1,tag2 atau path/ke/list.txt' if ind else 'tag1,tag2 or path/to/list.txt')
        print_format = QLineEdit()
        print_format.setPlaceholderText('id | {category}/{id} | post:{md5[:8]}')
        print_to_file = QLineEdit()
        print_to_file.setPlaceholderText('File output opsional untuk --print-to-file' if ind else 'Optional output file for --print-to-file')
        write_log = QLineEdit()
        write_log.setPlaceholderText('gallery-dl.log')
        error_file = QLineEdit()
        error_file.setPlaceholderText('failed_urls.txt')
        unsupported_file = QLineEdit()
        unsupported_file.setPlaceholderText('unsupported_urls.txt')
        config_json = QLineEdit()
        config_json.setPlaceholderText('Path file config JSON opsional untuk --config-json' if ind else 'Optional JSON config file path for --config-json')
        extra = QLineEdit()
        extra.setPlaceholderText("Argumen ekstra, contoh --write-pages --print-traffic" if ind else "Extra args, e.g. --write-pages --print-traffic")
        use_config = QCheckBox(bt("Use current GUI config (--config)", "Pakai config GUI saat ini (--config)"))
        use_config.setChecked(bool(self.config_path and Path(str(self.config_path)).expanduser().is_file()))
        simulate = QCheckBox(bt("Simulate only (-s)", "Simulasi saja (-s)"))
        get_urls = QCheckBox(bt("Get URLs (-g)", "Ambil URL (-g)"))
        dump_json = QCheckBox(bt("Dump JSON (-j)", "Dump JSON (-j)"))
        list_keywords = QCheckBox(bt("List keywords (-K)", "Lihat keyword (-K)"))
        extractor_info = QCheckBox(bt("Extractor info (-E)", "Info extractor (-E)"))
        write_metadata = QCheckBox(bt("Write metadata", "Tulis metadata"))
        write_info_json = QCheckBox(bt("Write info.json", "Tulis info.json"))
        zip_dl = QCheckBox(bt("gallery-dl --zip", "gallery-dl --zip"))
        cbz_dl = QCheckBox(bt("gallery-dl --cbz", "gallery-dl --cbz"))
        no_part = QCheckBox(bt("No .part files", "Tanpa file .part"))
        no_mtime = QCheckBox(bt("No mtime", "Tanpa mtime"))
        windows_names = QCheckBox(bt("Windows-safe filenames", "Nama file aman Windows"))
        no_cert = QCheckBox(bt("No certificate check / skip SSL cert", "Tanpa cek sertifikat / skip SSL"))
        no_skip = QCheckBox(bt("No skip / overwrite existing", "Tanpa skip / timpa file lama"))
        config_ignore = QCheckBox(bt("Ignore default config", "Abaikan config default"))
        netrc_auth = QCheckBox(bt("Use .netrc auth", "Pakai auth .netrc"))
        quiet_mode = QCheckBox(bt("Quiet", "Senyap"))
        warning_mode = QCheckBox(bt("Warnings only", "Peringatan saja"))
        verbose_mode = QCheckBox(bt("Verbose debug", "Debug rinci"))
        no_input = QCheckBox(bt("No input prompts", "Tanpa prompt input"))
        no_colors = QCheckBox(bt("No ANSI colors", "Tanpa warna ANSI"))
        write_pages = QCheckBox(bt("Write pages for debug", "Tulis halaman untuk debug"))
        print_traffic = QCheckBox(bt("Print HTTP traffic", "Cetak traffic HTTP"))
        resolve_urls = QCheckBox(bt("Resolve URLs (-G)", "Resolve URL (-G)"))
        resolve_json = QCheckBox(bt("Resolve JSON (-J)", "Resolve JSON (-J)"))
        no_download = QCheckBox(bt("No download", "Tanpa download"))

        fields = [
            ("Range", range_edit), ("Post range", post_range_edit),
            ("Child range", child_range_edit), ("Date after", date_after),
            ("Date before", date_before), ("Sleep", sleep_edit),
            ("Sleep request", sleep_request), ("Sleep 429", sleep_429),
            ("Sleep skip", sleep_skip), ("Sleep extractor", sleep_extractor),
            ("Sleep retries", sleep_retries), ("Retries", retries_spin),
            ("Limit rate", limit_rate), ("Chunk size", chunk_size),
            ("Min size", filesize_min), ("Max size", filesize_max),
            ("HTTP timeout", timeout_spin), ("Download archive", archive_path),
            ("Abort after skip", abort_after), ("Terminate after skip", terminate_after),
            ("Cookies browser", cookies_browser), ("Cookies file", cookies_file),
            ("Cookies export", cookies_export),
            ("Proxy", proxy), ("User-Agent", user_agent),
            ("-o options", option_kv), ("Filter", filter_expr),
            ("Post filter", post_filter_expr), ("Child filter", child_filter_expr),
            ("Blacklist", blacklist), ("Whitelist", whitelist),
            ("Tags blacklist", tags_blacklist), ("Tags whitelist", tags_whitelist),
            ("Print format", print_format), ("Print file", print_to_file),
            ("Write log", write_log), ("Error file", error_file),
            ("Unsupported file", unsupported_file), ("Config JSON file", config_json),
            ("Extra args", extra),
        ]
        field_label_id = {
            "Range": "Rentang", "Post range": "Rentang post", "Child range": "Rentang child",
            "Date after": "Tanggal setelah", "Date before": "Tanggal sebelum", "Sleep": "Jeda",
            "Sleep request": "Jeda request", "Sleep 429": "Jeda 429",
            "Sleep skip": "Jeda saat skip", "Sleep extractor": "Jeda extractor",
            "Sleep retries": "Jeda retry", "Retries": "Coba ulang",
            "Limit rate": "Batas kecepatan", "Chunk size": "Ukuran chunk",
            "Min size": "Ukuran min", "Max size": "Ukuran maks",
            "HTTP timeout": "Timeout HTTP", "Download archive": "Arsip download",
            "Abort after skip": "Abort setelah skip", "Terminate after skip": "Terminate setelah skip",
            "Cookies browser": "Cookies browser", "Cookies file": "File cookies", "Cookies export": "Ekspor cookies", "Proxy": "Proxy",
            "User-Agent": "User-Agent", "-o options": "Opsi -o", "Filter": "Filter",
            "Post filter": "Filter post", "Child filter": "Filter child", "Blacklist": "Daftar hitam",
            "Whitelist": "Daftar putih", "Tags blacklist": "Daftar hitam tag", "Tags whitelist": "Daftar putih tag",
            "Print format": "Format cetak", "Print file": "File cetak", "Write log": "Tulis log",
            "Error file": "File error", "Unsupported file": "File tak didukung",
            "Config JSON file": "File config JSON", "Extra args": "Argumen ekstra",
        }
        for i, (label_text, widget) in enumerate(fields):
            label_text = field_label_id.get(label_text, label_text) if ind else label_text
            r = i // 2
            c = (i % 2) * 2
            label = QLabel(label_text)
            label.setObjectName("fieldLabel")
            opt_l.addWidget(label, r, c)
            opt_l.addWidget(widget, r, c + 1)
        check_grid = QGridLayout()
        checks = [use_config, simulate, get_urls, dump_json, list_keywords, extractor_info, write_metadata, write_info_json, zip_dl, cbz_dl, no_part, no_mtime, windows_names, no_cert, no_skip, config_ignore, netrc_auth, quiet_mode, warning_mode, verbose_mode, no_input, no_colors, write_pages, print_traffic, resolve_urls, resolve_json, no_download]
        for i, chk in enumerate(checks):
            check_grid.addWidget(chk, i // 2, i % 2)
        opt_l.addLayout(check_grid, (len(fields)+1)//2, 0, 1, 4)
        opt_l.setColumnStretch(1, 1)
        opt_l.setColumnStretch(3, 1)

        # ----- Preview tab ----- #
        preview_page = QWidget()
        preview_l = QVBoxLayout(preview_page)
        preview_l.setContentsMargins(8, 8, 8, 8)
        tabs.addTab(preview_page, "Preview" if not ind else "Pratinjau")
        preview = QPlainTextEdit()
        preview.setReadOnly(True)
        preview.setMinimumHeight(180)
        preview_page_note = QLabel("Ini command tepat yang akan ditambahkan/diekspor. Quoting mengikuti OS saat ini." if ind else "This is the exact command that will be appended/exported. It is quoted for the current OS preview.")
        preview_page_note.setObjectName("subtle")
        preview_page_note.setWordWrap(True)
        preview_l.addWidget(preview_page_note)
        preview_l.addWidget(preview, 1)
        warn = QLabel("")
        warn.setObjectName("subtle")
        warn.setWordWrap(True)
        preview_l.addWidget(warn)

        # ----- Guide tab ----- #
        guide_page = QWidget()
        guide_l = QVBoxLayout(guide_page)
        guide_l.setContentsMargins(8, 8, 8, 8)
        tabs.addTab(guide_page, "Guide" if not ind else "Panduan")
        guide = QPlainTextEdit()
        guide.setReadOnly(True)
        guide.setPlainText(self.gallery_dl_command_guide_text("Indonesia" if ind else "English"))
        guide_l.addWidget(guide, 1)

        def add_opt(parts: list[str], flag: str, value: str) -> None:
            value = value.strip()
            if value:
                parts.extend([flag, value])

        def add_many_o_options(parts: list[str], raw: str) -> None:
            raw = raw.strip()
            if not raw:
                return
            chunks = [x.strip() for x in re.split(r"[;\n]+", raw) if x.strip()]
            for chunk in chunks:
                parts.extend(["-o", chunk])

        def current_parts() -> list[str]:
            parts = ["gallery-dl"]
            if use_config.isChecked() and self.config_path and Path(str(self.config_path)).expanduser().is_file():
                parts += ["--config", str(self.config_path)]
            add_opt(parts, "-d", dest.text())
            add_opt(parts, "-D", exact_dir.text())
            add_opt(parts, "-f", filename.text())
            add_opt(parts, "--range", range_edit.text())
            add_opt(parts, "--post-range", post_range_edit.text())
            add_opt(parts, "--child-range", child_range_edit.text())
            add_opt(parts, "--date-after", date_after.text())
            add_opt(parts, "--date-before", date_before.text())
            add_opt(parts, "--sleep", sleep_edit.text())
            add_opt(parts, "--sleep-request", sleep_request.text())
            if sleep_429.text().strip():
                parts += ["-o", f"sleep-429={sleep_429.text().strip()}"]
            add_opt(parts, "--sleep-skip", sleep_skip.text())
            add_opt(parts, "--sleep-extractor", sleep_extractor.text())
            if sleep_retries.text().strip():
                parts += ["-o", f"sleep-retries={sleep_retries.text().strip()}"]
            if retries_spin.value() > 0:
                parts += ["-R", str(retries_spin.value())]
            add_opt(parts, "--limit-rate", limit_rate.text())
            add_opt(parts, "--chunk-size", chunk_size.text())
            add_opt(parts, "--filesize-min", filesize_min.text())
            add_opt(parts, "--filesize-max", filesize_max.text())
            if timeout_spin.value() > 0:
                parts += ["--http-timeout", str(timeout_spin.value())]
            add_opt(parts, "--download-archive", archive_path.text())
            add_opt(parts, "-A", abort_after.text())
            add_opt(parts, "-T", terminate_after.text())
            browser_text = cookies_browser.currentText().strip()
            browser = dict(zip(cookies_items_id, cookies_items_en, strict=False)).get(browser_text, browser_text).lower()
            if browser == "use gui selection":
                browser = self.combo_cookies.currentText().strip().lower()
            if browser and browser != "none":
                parts += ["--cookies-from-browser", browser]
            add_opt(parts, "-C", cookies_file.text())
            add_opt(parts, "--cookies-export", cookies_export.text())
            add_opt(parts, "--proxy", proxy.text())
            add_opt(parts, "-a", user_agent.text())
            add_many_o_options(parts, option_kv.text())
            add_opt(parts, "--filter", filter_expr.text())
            add_opt(parts, "--post-filter", post_filter_expr.text())
            add_opt(parts, "--child-filter", child_filter_expr.text())
            add_opt(parts, "--blacklist", blacklist.text())
            add_opt(parts, "--whitelist", whitelist.text())
            add_opt(parts, "--tags-blacklist", tags_blacklist.text())
            add_opt(parts, "--tags-whitelist", tags_whitelist.text())
            if print_format.text().strip() and print_to_file.text().strip():
                parts.extend(["--print-to-file", print_format.text().strip(), print_to_file.text().strip()])
            else:
                add_opt(parts, "-N", print_format.text())
            add_opt(parts, "--write-log", write_log.text())
            add_opt(parts, "-e", error_file.text())
            add_opt(parts, "--write-unsupported", unsupported_file.text())
            add_opt(parts, "--config-json", config_json.text())
            if simulate.isChecked():
                parts.append("-s")
            if get_urls.isChecked():
                parts.append("-g")
            if dump_json.isChecked():
                parts.append("-j")
            if list_keywords.isChecked():
                parts.append("-K")
            if extractor_info.isChecked():
                parts.append("-E")
            if write_metadata.isChecked():
                parts.append("--write-metadata")
            if write_info_json.isChecked():
                parts.append("--write-info-json")
            if zip_dl.isChecked():
                parts.append("--zip")
            if cbz_dl.isChecked():
                parts.append("--cbz")
            if no_part.isChecked():
                parts.append("--no-part")
            if no_mtime.isChecked():
                parts.append("--no-mtime")
            if windows_names.isChecked():
                parts.append("--windows-filenames")
            if no_cert.isChecked():
                parts.append("--no-check-certificate")
            if no_skip.isChecked():
                parts.append("--no-skip")
            if config_ignore.isChecked():
                parts.append("--config-ignore")
            if netrc_auth.isChecked():
                parts.append("--netrc")
            if quiet_mode.isChecked():
                parts.append("-q")
            if warning_mode.isChecked():
                parts.append("-w")
            if verbose_mode.isChecked():
                parts.append("-v")
            if no_input.isChecked():
                parts.append("--no-input")
            if no_colors.isChecked():
                parts.append("--no-colors")
            if write_pages.isChecked():
                parts.append("--write-pages")
            if print_traffic.isChecked():
                parts.append("--print-traffic")
            if resolve_urls.isChecked():
                parts.append("-G")
            if resolve_json.isChecked():
                parts.append("-J")
            if no_download.isChecked():
                parts.append("--no-download")
            if extra.text().strip():
                parts += split_command(extra.text().strip())
            if url.text().strip():
                parts.append(url.text().strip())
            return parts

        def current_command() -> str:
            return " ".join(quote_arg_for_preview(x) for x in current_parts())

        def current_row() -> dict[str, str]:
            row = {
                "url": url.text().strip(),
                "destination": dest.text().strip(),
                "extra_args": extra.text().strip(),
                "enabled": "true",
                "tag": tag.text().strip(),
                "notes": notes.text().strip(),
                "command": current_command(),
            }
            # CSV/XLSX opened in spreadsheet apps must not execute formulas
            # supplied through a URL, tag, notes, destination, or extra args.
            return {key: sanitize_spreadsheet_cell(value) for key, value in row.items()}

        def ensure_url() -> bool:
            raw_url = url.text().strip()
            if not raw_url:
                self.show_compact_message("Command Builder", "URL is required." if not ind else "URL wajib diisi.", "warning")
                return False
            if any(ch in raw_url for ch in "\r\n"):
                self.show_compact_message("Command Builder", "URL must be a single line." if not ind else "URL harus satu baris.", "warning")
                return False
            if raw_url.startswith("-"):
                self.show_compact_message("Command Builder", "URL cannot start with an option flag." if not ind else "URL tidak boleh diawali flag opsi.", "warning")
                return False
            if " " in raw_url or "\t" in raw_url:
                res = QMessageBox.question(
                    dlg,
                    "URL contains whitespace" if not ind else "URL berisi spasi",
                    (
                        "The URL contains whitespace. The builder will quote it as one argument. Continue?"
                        if not ind else
                        "URL berisi spasi. Builder akan mengutipnya sebagai satu argumen. Lanjutkan?"
                    ),
                )
                if res != QMessageBox.Yes:
                    return False
            if not (raw_url.startswith(("http://", "https://")) or re.match(r"^[A-Za-z0-9_.+-]+:https?://", raw_url) or raw_url.startswith(("r:http://", "r:https://"))):
                res = QMessageBox.question(
                    dlg,
                    "Unusual URL format" if not ind else "Format URL tidak biasa",
                    (
                        "This does not look like a normal URL or extractor-prefixed URL. gallery-dl may still support it. Continue?"
                        if not ind else
                        "Ini tidak terlihat seperti URL normal atau URL dengan prefix extractor. gallery-dl mungkin tetap mendukungnya. Lanjutkan?"
                    ),
                )
                if res != QMessageBox.Yes:
                    return False
            return True

        def refresh_preview() -> None:
            cmd = current_command()
            preview.setPlainText(cmd)
            issues: list[str] = []
            if not url.text().strip():
                issues.append("URL is still empty." if not ind else "URL masih kosong.")
            if dest.text().strip() and exact_dir.text().strip():
                issues.append("Both -d and -D are set. Use only one unless you know exactly why." if not ind else "-d dan -D sama-sama aktif. Gunakan salah satu kecuali memang sengaja.")
            if get_urls.isChecked() and (write_metadata.isChecked() or zip_dl.isChecked() or cbz_dl.isChecked()):
                issues.append("-g prints URLs instead of downloading, so post-processing options may not do what you expect." if not ind else "-g hanya mencetak URL, jadi opsi post-processing mungkin tidak berguna.")
            if zip_dl.isChecked() and cbz_dl.isChecked():
                issues.append("Both --zip and --cbz are enabled. Usually choose one archive type." if not ind else "--zip dan --cbz sama-sama aktif. Biasanya pilih salah satu format arsip.")
            if quiet_mode.isChecked() and verbose_mode.isChecked():
                issues.append("Quiet and verbose conflict conceptually. Choose one output mode." if not ind else "Quiet dan verbose bertabrakan secara konsep. Pilih salah satu mode output.")
            if print_format.text().strip() and not print_to_file.text().strip():
                issues.append("-N/--print writes to standard output. Use Print file if you want a file." if not ind else "-N/--print menulis ke standard output. Isi Print file jika ingin disimpan ke file.")
            warn.setText("\n".join("[WARNING] " + x for x in issues))

        def apply_preset(name: str) -> None:
            name = preset_key(name)
            # Clear only option widgets, not URL/tag/notes.
            for w in (exact_dir, filename, range_edit, post_range_edit, child_range_edit, date_after, date_before, sleep_edit, sleep_request, sleep_429, sleep_skip, sleep_extractor, sleep_retries, limit_rate, chunk_size, filesize_min, filesize_max, archive_path, abort_after, terminate_after, cookies_file, cookies_export, proxy, user_agent, option_kv, filter_expr, post_filter_expr, child_filter_expr, blacklist, whitelist, tags_blacklist, tags_whitelist, print_format, print_to_file, write_log, error_file, unsupported_file, config_json, extra):
                w.clear()
            retries_spin.setValue(0)
            timeout_spin.setValue(0)
            # Reset by index: setCurrentText("none") is a silent no-op in the
            # Indonesian UI where item 0 is "tidak ada", so switching presets
            # kept a stale browser-cookies choice.
            cookies_browser.setCurrentIndex(0)
            for chk in (simulate, get_urls, dump_json, list_keywords, extractor_info, write_metadata, write_info_json, zip_dl, cbz_dl, no_part, no_mtime, windows_names, no_cert, no_skip, config_ignore, netrc_auth, quiet_mode, warning_mode, verbose_mode, no_input, no_colors, write_pages, print_traffic, resolve_urls, resolve_json, no_download):
                chk.setChecked(False)
            if name == "Download to selected folder":
                if not dest.text().strip():
                    dest.setText(self.edit_output.text().strip() or "./downloads")
            elif name == "Safe slow mode":
                sleep_edit.setText("2.0-5.0")
                sleep_request.setText("1.0")
                sleep_429.setText("60-300")
                retries_spin.setValue(max(4, self.spin_retries.value()))
            elif name == "Preview only / simulate":
                simulate.setChecked(True)
            elif name == "Get direct media URLs":
                get_urls.setChecked(True)
            elif name == "Range 1-20":
                range_edit.setText("1-20")
            elif name == "Use browser cookies":
                cookies_browser.setCurrentText(("pakai pilihan GUI" if ind else "use GUI selection") if self.combo_cookies.currentText() != "none" else "firefox")
            elif name == "Archive skip downloaded":
                archive_path.setText(str(APP_DIR / "gallery-dl-archive.sqlite3"))
            elif name == "Metadata + info JSON":
                write_metadata.setChecked(True)
                write_info_json.setChecked(True)
            elif name == "Windows-safe filenames":
                windows_names.setChecked(True)
            elif name == "Keyword discovery (-K)":
                list_keywords.setChecked(True)
                simulate.setChecked(True)
            elif name == "Extractor debug info (-E)":
                extractor_info.setChecked(True)
                verbose_mode.setChecked(True)
            elif name == "JSON metadata audit":
                dump_json.setChecked(True)
                write_info_json.setChecked(True)
            elif name == "Rate-limit friendly":
                sleep_edit.setText("3.0-7.0")
                sleep_request.setText("1.0-2.5")
                sleep_429.setText("120-600")
                retries_spin.setValue(max(6, self.spin_retries.value()))
            elif name == "Write logs + error file":
                write_log.setText(str(APP_DIR / "gallery-dl-run.log"))
                error_file.setText(str(APP_DIR / "gallery-dl-errors.txt"))
                unsupported_file.setText(str(APP_DIR / "gallery-dl-unsupported.txt"))
            elif name == "ZIP archive output":
                zip_dl.setChecked(True)
            elif name == "CBZ manga archive":
                cbz_dl.setChecked(True)
            elif name == "Date filter example":
                date_after.setText("2026-01-01")
                date_before.setText("2026-12-31")
            elif name == "Tag whitelist example":
                tags_whitelist.setText("highres,illustration")
            elif name == "File size filter example":
                filesize_min.setText("100k")
                filesize_max.setText("50M")
            elif name == "SSL certificate bypass":
                no_cert.setChecked(True)
                retries_spin.setValue(max(4, self.spin_retries.value()))
                timeout_spin.setValue(60)
                verbose_mode.setChecked(True)
            elif name == "HypnoHub cautious mode":
                if not url.text().strip():
                    url.setText("https://hypnohub.net/index.php?page=post&s=list&tags=TAG")
                sleep_request.setText("2.0-5.0")
                sleep_429.setText("120-600")
                sleep_retries.setText("60-300")
                retries_spin.setValue(max(6, self.spin_retries.value()))
                archive_path.setText(str(APP_DIR / "hypnohub-archive.sqlite3"))
                terminate_after.setText("5")
                windows_names.setChecked(True)
                no_input.setChecked(True)
                filename.setText("{id}_{md5}.{extension}")
            elif name == "Rule34 cautious mode":
                if not url.text().strip():
                    url.setText("https://rule34.xxx/index.php?page=post&s=list&tags=TAG")
                sleep_request.setText("2.0-4.0")
                sleep_429.setText("120-600")
                retries_spin.setValue(max(6, self.spin_retries.value()))
                archive_path.setText(str(APP_DIR / "rule34-archive.sqlite3"))
                terminate_after.setText("5")
                windows_names.setChecked(True)
                no_input.setChecked(True)
            elif name == "Booru tag batch":
                sleep_request.setText("1.5-3.0")
                sleep_skip.setText("0.2-1.0")
                terminate_after.setText("8")
                archive_path.setText(str(APP_DIR / "booru-archive.sqlite3"))
                windows_names.setChecked(True)
                tags_blacklist.setText("animated_gif,lowres")
            elif name == "Resume large batch":
                archive_path.setText(str(APP_DIR / "gallery-dl-archive.sqlite3"))
                retries_spin.setValue(max(8, self.spin_retries.value()))
                sleep_retries.setText("60-300")
                terminate_after.setText("10")
                no_part.setChecked(False)
            refresh_preview()

        widgets = [url, dest, exact_dir, filename, tag, notes, range_edit, post_range_edit, child_range_edit, date_after, date_before, sleep_edit, sleep_request, sleep_429, limit_rate, filesize_min, filesize_max, archive_path, cookies_file, proxy, user_agent, option_kv, filter_expr, post_filter_expr, child_filter_expr, blacklist, whitelist, tags_blacklist, tags_whitelist, print_format, print_to_file, write_log, error_file, unsupported_file, config_json, extra]
        for widget in widgets:
            widget.textChanged.connect(refresh_preview)
        for widget in (retries_spin, timeout_spin):
            widget.valueChanged.connect(refresh_preview)
        for widget in (cookies_browser,):
            widget.currentTextChanged.connect(refresh_preview)
        for chk in (use_config, simulate, get_urls, dump_json, list_keywords, extractor_info, write_metadata, write_info_json, zip_dl, cbz_dl, no_part, no_mtime, windows_names, no_cert, no_skip, config_ignore, netrc_auth, quiet_mode, warning_mode, verbose_mode, no_input, no_colors, write_pages, print_traffic, resolve_urls, resolve_json, no_download):
            chk.stateChanged.connect(refresh_preview)
        preset.currentTextChanged.connect(apply_preset)

        # ----- Buttons ----- #
        btn_row = QHBoxLayout()
        add_btn = QPushButton("Append to Input" if not ind else "Tambah ke Input")
        copy_btn = QPushButton("Copy Command" if not ind else "Salin Command")
        txt_btn = QPushButton("Save TXT" if not ind else "Simpan TXT")
        csv_btn = QPushButton("Save CSV" if not ind else "Simpan CSV")
        xlsx_btn = QPushButton("Save XLSX" if not ind else "Simpan XLSX")
        guide_copy_btn = QPushButton("Copy Guide" if not ind else "Salin Panduan")
        for b in (add_btn, copy_btn, txt_btn, csv_btn, xlsx_btn, guide_copy_btn):
            b.setMinimumHeight(30)
        btn_row.addWidget(add_btn)
        btn_row.addWidget(copy_btn)
        btn_row.addWidget(txt_btn)
        btn_row.addWidget(csv_btn)
        btn_row.addWidget(xlsx_btn)
        btn_row.addWidget(guide_copy_btn)
        btn_row.addStretch(1)
        root.addLayout(btn_row)

        def append_cmd() -> None:
            if not ensure_url():
                return
            cmd = current_command()
            current = self.txt_commands.toPlainText().strip()
            self.txt_commands.setPlainText((current + "\n" + cmd).strip() if current else cmd)
            self.append_log("[builder] command appended to main input")

        def copy_cmd() -> None:
            QApplication.clipboard().setText(current_command())
            self.show_compact_message("Command Builder", "Command copied." if not ind else "Command disalin.", "info")

        def append_txt() -> None:
            if not ensure_url():
                return
            path, _ = QFileDialog.getSaveFileName(dlg, "Save TXT database", "gallery_dl_database.txt", "Text (*.txt);;All files (*.*)")
            if not path:
                return
            try:
                p = Path(path)
                existed = p.exists() and p.stat().st_size > 0
                p.parent.mkdir(parents=True, exist_ok=True)
                with p.open("a", encoding="utf-8", newline="") as f:
                    if existed:
                        f.write("\n")
                    if tag.text().strip():
                        f.write(f"# {tag.text().strip()}\n")
                    f.write(current_command() + "\n")
                self.show_compact_message("Command Builder", f"TXT saved:\n{p}", "info")
            except Exception as exc:
                self.show_compact_message("Save TXT failed", str(exc), "error")

        def append_csv() -> None:
            if not ensure_url():
                return
            path, _ = QFileDialog.getSaveFileName(dlg, "Save CSV database", "gallery_dl_database.csv", "CSV (*.csv);;All files (*.*)")
            if not path:
                return
            try:
                p = Path(path)
                p.parent.mkdir(parents=True, exist_ok=True)
                new_file = (not p.exists()) or p.stat().st_size == 0
                headers = ["url", "destination", "extra_args", "enabled", "tag", "notes", "command"]
                with p.open("a", encoding="utf-8-sig", newline="") as f:
                    writer = csv.DictWriter(f, fieldnames=headers)
                    if new_file:
                        writer.writeheader()
                    writer.writerow(current_row())
                self.show_compact_message("Command Builder", f"CSV saved:\n{p}", "info")
            except Exception as exc:
                self.show_compact_message("Save CSV failed", str(exc), "error")

        def append_xlsx() -> None:
            if not ensure_url():
                return
            try:
                import openpyxl  # type: ignore
            except Exception:
                self.show_compact_message("Command Builder", "openpyxl is required for XLSX export.\n\nInstall with:\npip install openpyxl", "warning")
                return
            path, _ = QFileDialog.getSaveFileName(dlg, "Save XLSX database", "gallery_dl_database.xlsx", "Excel (*.xlsx);;All files (*.*)")
            if not path:
                return
            wb = None
            try:
                p = Path(path)
                p.parent.mkdir(parents=True, exist_ok=True)
                headers = ["url", "destination", "extra_args", "enabled", "tag", "notes", "command"]
                if p.exists() and p.stat().st_size > 0:
                    # Same size guard as the importer (stab3.3): appending to a
                    # user-chosen huge workbook would otherwise load it whole.
                    if p.stat().st_size > MAX_IMPORT_BYTES:
                        self.show_compact_message(
                            "Save XLSX failed",
                            f"Existing XLSX is too large to append to "
                            f"({p.stat().st_size // (1024 * 1024)} MB > {MAX_IMPORT_BYTES // (1024 * 1024)} MB limit).",
                            "error",
                        )
                        return
                    wb = openpyxl.load_workbook(p)
                    ws = wb.active
                    # openpyxl reports max_row == 1 even for an empty sheet, so
                    # "max_row == 0" never fired and an empty active sheet got
                    # data rows with no header (misimported later). Check the
                    # first row's contents instead.
                    first_row_empty = all(cell.value in (None, "") for cell in ws[1])
                    if first_row_empty:
                        for col_i, header in enumerate(headers, start=1):
                            ws.cell(row=1, column=col_i, value=header)
                else:
                    wb = openpyxl.Workbook()
                    ws = wb.active
                    ws.title = "gallery-dl database"
                    ws.append(headers)
                row = current_row()
                ws.append([row[h] for h in headers])
                # Force the appended data cells to Text so values that begin
                # with "-"/"--" (extra_args, option-heavy commands) are not
                # misread by Excel as formulas (#NAME?) — same protection the
                # template exporter applies to its data rows.
                for cell in ws[ws.max_row]:
                    cell.number_format = "@"
                for col in range(1, len(headers) + 1):
                    ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 24 if col != 7 else 90
                self._xlsx_save_atomic(wb, p)
                self.show_compact_message("Command Builder", f"XLSX saved:\n{p}", "info")
            except Exception as exc:
                self.show_compact_message("Save XLSX failed", str(exc), "error")
            finally:
                if wb is not None:
                    try:
                        wb.close()
                    except Exception:
                        pass

        def copy_guide() -> None:
            QApplication.clipboard().setText(self.gallery_dl_command_guide_text("Indonesia" if ind else "English"))
            self.show_compact_message("Command Builder", "Guide copied." if not ind else "Panduan disalin.", "info")

        add_btn.clicked.connect(append_cmd)
        copy_btn.clicked.connect(copy_cmd)
        txt_btn.clicked.connect(append_txt)
        csv_btn.clicked.connect(append_csv)
        xlsx_btn.clicked.connect(append_xlsx)
        guide_copy_btn.clicked.connect(copy_guide)
        refresh_preview()
        dlg.exec()

    def gallery_dl_config_template_text(self, preset: str = "Safe default") -> str:
        """Return a valid JSON config template for gallery-dl Builder presets."""
        archive = str(APP_DIR / "gallery-dl-archive.sqlite3").replace("\\", "/")
        out_dir = str(safe_expand_path(self.edit_output.text() or "./downloads")).replace("\\", "/")
        base = {
            "extractor": {
                "base-directory": out_dir,
                "directory": ["{category}", "{subcategory}", "{user[id]|username|author|id}"],
                "filename": "{id}_{filename}.{extension}",
                "archive": archive,
                "skip": "abort:3",
                "retries": max(4, int(self.spin_retries.value() or 4)),
                "timeout": 30.0,
                "sleep-request": "1.0-2.0",
                "sleep-429": "120-600",
                "path-restrict": "auto",
                "path-replace": "_",
                "path-strip": "auto"
            },
            "downloader": {
                "part": True,
                "mtime": True,
                "retries": max(4, int(self.spin_retries.value() or 4)),
                "timeout": 30.0,
                "rate": None
            },
            "output": {
                "progress": True,
                "log": "[{name}][{levelname}] {message}",
                "errorfile": str(APP_DIR / "gallery-dl-errors.txt").replace("\\", "/"),
                "unsupportedfile": str(APP_DIR / "gallery-dl-unsupported.txt").replace("\\", "/")
            }
        }
        key = str(preset or "Safe default").lower()
        if "archive" in key:
            base["extractor"].update({
                "archive": archive,
                "archive-event": "file",
                "archive-mode": "file",
                "archive-pragma": ["journal_mode=WAL", "synchronous=NORMAL"],
                "skip": "abort:5"
            })
        elif "cookies" in key:
            browser = self.combo_cookies.currentText().strip().lower() or "firefox"
            if browser == "none":
                browser = "firefox"
            base["extractor"].update({
                "cookies": [browser],
                "cookies-update": True,
                "parent-session": True
            })
        elif "rate" in key:
            base["extractor"].update({
                "sleep": "2.0-5.0",
                "sleep-request": "1.5-3.0",
                "sleep-retries": "60-300",
                "sleep-429": "120-900",
                "retries": 8,
                "retry-codes": [429, 430, 500, 502, 503, 504]
            })
            base["downloader"].update({"rate": "800k-2M", "retries": 8})
        elif "windows" in key:
            base["extractor"].update({
                "path-restrict": "windows",
                "path-strip": "windows",
                "path-replace": "_",
                "path-extended": True,
                "filename": "{id}_{title|filename}.{extension}"
            })
        elif "metadata" in key:
            base["extractor"].update({
                "postprocessors": [
                    {"name": "metadata", "mode": "json"},
                    {"name": "metadata", "mode": "jsonl", "extension": "jsonl"}
                ]
            })
            base["output"].update({"private": True, "num-to-str": False})
        elif "site override" in key:
            base["extractor"].update({
                "pixiv": {
                    "directory": ["pixiv", "{user[id]} - {user[name]}"],
                    "filename": "{id}_p{num}.{extension}",
                    "archive": [":~", ".gallery-dl", "pixiv.sqlite3"],
                    "sleep-request": "1.0-2.5"
                },
                "twitter": {
                    "directory": ["twitter", "{user[name]}"],
                    "filename": "{tweet_id}_{num}.{extension}",
                    "sleep-request": "2.0-4.0"
                }
            })
        elif "debug" in key:
            base["output"].update({
                "mode": "terminal",
                "progress": True,
                "log": {"level": "debug", "format": "[{name}][{levelname}] {message}"},
                "logfile": str(APP_DIR / "gallery-dl-debug.log").replace("\\", "/")
            })
            base["extractor"].update({"write-pages": True, "retries": 1})
        return json.dumps(base, indent=4, ensure_ascii=False)

    def gallery_dl_command_guide_text(self, language: str | None = None) -> str:
        language = language or getattr(self, "help_language", "English")
        if language == "Indonesia":
            return """Panduan Lengkap Builder gallery-dl
=================================

Sumber resmi:
- Project: https://codeberg.org/mikf/gallery-dl
- Opsi command-line: https://gdl-org.github.io/docs/options.html
- Konfigurasi JSON: https://gdl-org.github.io/docs/configuration.html

1) Pola dasar
-------------
  gallery-dl [OPTIONS]... URLS...

Builder ini hanya menyusun command. Download tetap dijalankan oleh aplikasi utama melalui subprocess gallery-dl. Jadi command yang dibuat di Builder aman untuk ditinjau, disalin, disimpan ke TXT/CSV/XLSX, atau ditambahkan ke input utama.

2) Output folder: -d vs -D vs -f
--------------------------------
-d / --destination PATH
  Folder dasar output. Cocok untuk workflow umum.

-D / --directory PATH
  Folder exact. Gunakan jika ingin hasil masuk tepat ke folder itu.

-f / --filename FORMAT
  Format nama file. Gunakan keyword dari extractor. Cara mencari keyword:
    gallery-dl -K URL

Contoh:
  gallery-dl -d "D:\\Rips" -f "{id}_{filename}.{extension}" URL

3) Anti download ulang: archive
-------------------------------
Pakai salah satu:
  gallery-dl --download-archive "D:\\Rips\\archive.sqlite3" URL

atau di config:
  {
      "extractor": {
          "archive": "D:/Rips/archive.sqlite3",
          "skip": "abort:5"
      }
  }

Catatan:
- Archive gallery-dl adalah database SQLite, bukan TXT biasa.
- File archive otomatis dibuat jika belum ada.
- Archive berbeda dari session GUI.

4) Cookies dan login
--------------------
Command cepat:
  gallery-dl --cookies-from-browser firefox URL
  gallery-dl --cookies-from-browser chrome URL
  gallery-dl -C cookies.txt URL

Config browser cookies:
  {
      "extractor": {
          "cookies": ["firefox"],
          "cookies-update": true
      }
  }

Config cookies.txt:
  {
      "extractor": {
          "cookies": "D:/cookies/site-cookies.txt"
      }
  }

Tips:
- Gunakan cookies jika situs butuh login, 403, forbidden, atau konten tidak muncul.
- Jangan bagikan file cookies ke orang lain.

5) Rate-limit dan stabilitas jaringan
-------------------------------------
Command:
  gallery-dl --sleep 2.0-5.0 --sleep-request 1.0-2.0 -o sleep-429=120-600 -R 8 URL

Config:
  {
      "extractor": {
          "sleep": "2.0-5.0",
          "sleep-request": "1.0-2.0",
          "sleep-retries": "60-300",
          "sleep-429": "120-600",
          "retries": 8,
          "timeout": 30.0
      },
      "downloader": {
          "rate": "800k-2M",
          "retries": 8,
          "timeout": 30.0
      }
  }

Gunakan preset ini jika sering kena 429, timeout, temporary failure, atau koneksi putus.

6) Selection filter
-------------------
Range:
  --range 1-20

Tanggal (key config khusus extractor; dukungan tergantung situs):
  -o date-min=2026-01-01
  -o date-max=2026-12-31

Ukuran:
  --filesize-min 100k
  --filesize-max 50M

Filter Python expression:
  --filter "image_width >= 1000 and extension in ('jpg', 'png')"

Tag/metadata situs memakai expression --filter setelah mengecek keyword dengan -K.

7) Output audit dan debug
-------------------------
Tidak download, hanya uji ekstraksi:
  gallery-dl -s URL

Cetak direct media URL:
  gallery-dl -g URL

Cetak metadata JSON:
  gallery-dl -j URL

Lihat keyword:
  gallery-dl -K URL

Info extractor:
  gallery-dl -E URL

Tulis log:
  gallery-dl --write-log gallery-dl.log -e failed.txt --write-unsupported unsupported.txt URL

8) Config file lengkap
----------------------
File config gallery-dl memakai format JSON. Struktur penting:

  {
      "extractor": {
          "base-directory": "D:/Rips/gallery-dl",
          "directory": ["{category}", "{subcategory}", "{user[id]|id}"],
          "filename": "{id}_{filename}.{extension}",
          "archive": "D:/Rips/archive.sqlite3",
          "skip": "abort:5",
          "cookies": ["firefox"],
          "sleep-request": "1.0-2.0",
          "sleep-429": "120-600",
          "retries": 8,
          "timeout": 30.0,
          "path-restrict": "auto",
          "path-replace": "_",
          "path-strip": "auto"
      },
      "downloader": {
          "part": true,
          "mtime": true,
          "rate": "800k-2M",
          "retries": 8,
          "timeout": 30.0
      },
      "output": {
          "progress": true,
          "log": "[{name}][{levelname}] {message}",
          "errorfile": "D:/Rips/errors.txt",
          "unsupportedfile": "D:/Rips/unsupported.txt"
      }
  }

Cara pakai config:
  gallery-dl --config "D:\\Rips\\gallery-dl-config.json" URL

atau lewat field System > Config di GUI.

PENTING:
- --config-json menerima PATH FILE, bukan isi JSON langsung.
- Untuk override cepat per command, gunakan -o KEY=VALUE.
- Contoh -o:
    gallery-dl -o extractor.sleep-request=2.0 -o downloader.rate=1M URL

9) Level config: global, category, subcategory
----------------------------------------------
Config bisa ditaruh di beberapa level:

Base/global extractor:
  extractor.filename

Per situs/category:
  extractor.pixiv.filename

Per subcategory:
  extractor.pixiv.user.filename

Contoh:
  {
      "extractor": {
          "filename": "{id}.{extension}",
          "pixiv": {
              "directory": ["pixiv", "{user[id]}"],
              "filename": "{id}_p{num}.{extension}"
          },
          "twitter": {
              "directory": ["twitter", "{user[name]}"],
              "filename": "{tweet_id}_{num}.{extension}"
          }
      }
  }

10) Preset yang disarankan
--------------------------
Awal aman:
  - Gunakan archive.
  - Gunakan sleep-request.
  - Gunakan cookies bila situs butuh login.
  - Mulai dari 1 sampai 3 worker.
  - Jalankan Command Preview sebelum Start.

Untuk debug:
  - Coba -s, -K, -E, atau -j.
  - Baca error log.
  - Periksa cookies/config/path.

11) Penyebab gagal paling umum
------------------------------
- URL tidak didukung extractor gallery-dl.
- Cookies belum dipasang untuk situs login-only.
- Config JSON tidak valid.
- Path Windows mengandung karakter terlarang.
- Tidak memakai archive, sehingga file lama dicek ulang.
- Terlalu agresif, memicu 429/rate-limit.
"""
        return """Complete gallery-dl Builder Guide
=================================

Official sources:
- Project: https://codeberg.org/mikf/gallery-dl
- Command-line options: https://gdl-org.github.io/docs/options.html
- JSON configuration: https://gdl-org.github.io/docs/configuration.html

1) Basic pattern
----------------
  gallery-dl [OPTIONS]... URLS...

This Builder only creates commands. Downloads are still executed by the main app through a gallery-dl subprocess. You can review, copy, save, or append the generated command.

2) Output folder: -d vs -D vs -f
--------------------------------
-d / --destination PATH
  Base output folder. Best for normal use.

-D / --directory PATH
  Exact directory. Use when every file must go into one exact folder.

-f / --filename FORMAT
  Filename format. Use extractor keywords. Discover them with:
    gallery-dl -K URL

Example:
  gallery-dl -d "D:\\Rips" -f "{id}_{filename}.{extension}" URL

3) Avoid duplicate downloads: archive
-------------------------------------
Command:
  gallery-dl --download-archive "D:\\Rips\\archive.sqlite3" URL

Config:
  {
      "extractor": {
          "archive": "D:/Rips/archive.sqlite3",
          "skip": "abort:5"
      }
  }

Notes:
- gallery-dl archive is an SQLite database, not a plain TXT file.
- Missing archive files are created automatically.
- The gallery-dl archive is separate from GUI sessions.

4) Cookies and login
--------------------
Quick command:
  gallery-dl --cookies-from-browser firefox URL
  gallery-dl --cookies-from-browser chrome URL
  gallery-dl -C cookies.txt URL

Browser cookies config:
  {
      "extractor": {
          "cookies": ["firefox"],
          "cookies-update": true
      }
  }

cookies.txt config:
  {
      "extractor": {
          "cookies": "D:/cookies/site-cookies.txt"
      }
  }

Tips:
- Use cookies for login-only pages, 403, forbidden, or missing content.
- Never share cookie files.

5) Rate-limit and network stability
-----------------------------------
Command:
  gallery-dl --sleep 2.0-5.0 --sleep-request 1.0-2.0 -o sleep-429=120-600 -R 8 URL

Config:
  {
      "extractor": {
          "sleep": "2.0-5.0",
          "sleep-request": "1.0-2.0",
          "sleep-retries": "60-300",
          "sleep-429": "120-600",
          "retries": 8,
          "timeout": 30.0
      },
      "downloader": {
          "rate": "800k-2M",
          "retries": 8,
          "timeout": 30.0
      }
  }

Use this for 429, timeouts, temporary failures, or unstable networks.

6) Selection filters
--------------------
Ranges:
  --range 1-20

Dates (extractor config keys; support is site-specific):
  -o date-min=2026-01-01
  -o date-max=2026-12-31

File sizes:
  --filesize-min 100k
  --filesize-max 50M

Python expression filter:
  --filter "image_width >= 1000 and extension in ('jpg', 'png')"

Use --filter for site tag/metadata expressions after inspecting keywords with -K.

7) Audit and debug output
-------------------------
Simulate extraction:
  gallery-dl -s URL

Print direct media URLs:
  gallery-dl -g URL

Print JSON metadata:
  gallery-dl -j URL

List keywords:
  gallery-dl -K URL

Extractor defaults/settings:
  gallery-dl -E URL

Write logs:
  gallery-dl --write-log gallery-dl.log -e failed.txt --write-unsupported unsupported.txt URL

8) Complete config file
-----------------------
gallery-dl config files use JSON. Important structure:

  {
      "extractor": {
          "base-directory": "D:/Rips/gallery-dl",
          "directory": ["{category}", "{subcategory}", "{user[id]|id}"],
          "filename": "{id}_{filename}.{extension}",
          "archive": "D:/Rips/archive.sqlite3",
          "skip": "abort:5",
          "cookies": ["firefox"],
          "sleep-request": "1.0-2.0",
          "sleep-429": "120-600",
          "retries": 8,
          "timeout": 30.0,
          "path-restrict": "auto",
          "path-replace": "_",
          "path-strip": "auto"
      },
      "downloader": {
          "part": true,
          "mtime": true,
          "rate": "800k-2M",
          "retries": 8,
          "timeout": 30.0
      },
      "output": {
          "progress": true,
          "log": "[{name}][{levelname}] {message}",
          "errorfile": "D:/Rips/errors.txt",
          "unsupportedfile": "D:/Rips/unsupported.txt"
      }
  }

Use config:
  gallery-dl --config "D:\\Rips\\gallery-dl-config.json" URL

or use System > Config in this GUI.

IMPORTANT:
- --config-json expects a FILE PATH, not inline JSON text.
- For quick per-command overrides, use -o KEY=VALUE.
- Example:
    gallery-dl -o extractor.sleep-request=2.0 -o downloader.rate=1M URL

9) Config levels: global, category, subcategory
-----------------------------------------------
You can set options at multiple levels:

Base/global extractor:
  extractor.filename

Per site/category:
  extractor.pixiv.filename

Per subcategory:
  extractor.pixiv.user.filename

Example:
  {
      "extractor": {
          "filename": "{id}.{extension}",
          "pixiv": {
              "directory": ["pixiv", "{user[id]}"],
              "filename": "{id}_p{num}.{extension}"
          },
          "twitter": {
              "directory": ["twitter", "{user[name]}"],
              "filename": "{tweet_id}_{num}.{extension}"
          }
      }
  }

10) Recommended presets
-----------------------
Safe start:
  - Use archive.
  - Use sleep-request.
  - Use cookies when a site requires login.
  - Start with 1 to 3 workers.
  - Run Command Preview before Start.

Debug:
  - Try -s, -K, -E, or -j.
  - Read the error log.
  - Check cookies/config/path.

11) Common failure causes
-------------------------
- URL is not supported by gallery-dl.
- Login-only pages are used without cookies.
- JSON config is invalid.
- Windows path contains forbidden characters.
- --download-archive is not used, so old files are checked again.
- Requests are too aggressive and trigger 429/rate limits.
"""

    def generate_config_template(self) -> None:
        default = safe_expand_path(self.config_path or detect_config_path() or (APP_DIR / "gallery-dl-config-template.json"))
        title = "Simpan template config" if self._ui_is_indonesian() else "Save config template"
        path, _ = QFileDialog.getSaveFileName(self, title, str(default), "JSON (*.json);;All files (*.*)")
        if not path:
            return
        try:
            template = json.loads(self.gallery_dl_config_template_text("Safe default"))
            Path(path).parent.mkdir(parents=True, exist_ok=True)
            atomic_write_text(path, json.dumps(template, indent=4, ensure_ascii=False) + "\n", encoding="utf-8")
            self.show_compact_message("Config Template", f"Template saved:\n{path}", "info")
        except Exception as exc:
            self.show_compact_message("Config Template failed", str(exc), "error")

    def resume_unfinished_queue(self) -> None:
        if self.active_workers > 0:
            self.show_compact_message("Resume", "Wait until the current run finishes before resuming unfinished jobs.", "info")
            return
        # Honor the per-row enabled flag, matching Start's current_run_indices().
        # Without the filter, Resume queued rows the user explicitly disabled.
        idxs = [i for i, job in enumerate(self.jobs) if job.enabled and i not in self.done_indices]
        if not idxs:
            self.show_compact_message("Resume", "No unfinished jobs found.", "info")
            return
        self._start_indices(idxs, reset_result_sets=False)

    def save_project_profile(self) -> None:
        app_data_dir()
        path, _ = QFileDialog.getSaveFileName(self, "Save project profile", str(PROFILES_DIR / f"profile_{timestamp_slug()}.json"), "JSON (*.json)")
        if not path:
            return
        try:
            data = self.session_data()
            data["profile_saved_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
            atomic_write_text(path, json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
            self.show_compact_message("Project Profile", f"Profile saved:\n{path}", "info")
        except Exception as exc:
            self.show_compact_message("Save Project Profile failed", str(exc), "error")

    def load_project_profile(self) -> None:
        if self.active_workers > 0:
            self.show_compact_message("Project Profile", "Stop the current download before loading a project profile.", "info")
            return
        app_data_dir()
        path, _ = QFileDialog.getOpenFileName(self, "Load project profile", str(PROFILES_DIR), "JSON (*.json);;All files (*.*)")
        if not path:
            return
        try:
            data = json.loads(read_text_safely(path))
            if not isinstance(data, dict):
                raise ValueError("Profile must contain a JSON object.")
            self.apply_session_data(data)
            self.show_compact_message("Project Profile", f"Profile loaded:\n{path}", "info")
        except Exception as exc:
            self.show_compact_message("Load Project Profile failed", str(exc), "error")

    def set_compact_mode(self, enabled: bool, notify: bool = False) -> None:
        self.compact_mode = bool(enabled)
        if hasattr(self, "lbl_subtitle"):
            self.lbl_subtitle.setVisible(not self.compact_mode)
        if hasattr(self, "lbl_feature_hint"):
            self.lbl_feature_hint.setVisible(not self.compact_mode)
        if hasattr(self, "side_panel"):
            self.side_panel.setMinimumWidth(330 if self.compact_mode else 390)
            self.side_panel.setMaximumWidth(440 if self.compact_mode else 520)
        # Do not force-resize a maximized/fullscreen window — on Windows that
        # kicks the window out of the maximized state unexpectedly.
        if not (self.isMaximized() or self.isFullScreen()):
            self.resize(1080, 700) if self.compact_mode else self.resize(1220, 760)
        if notify:
            self.show_compact_message("Compact Mode", f"Compact mode is now {'ON' if self.compact_mode else 'OFF'}.", "info")

    def toggle_compact_mode(self) -> None:
        self.set_compact_mode(not self.compact_mode, notify=True)
        self.autosave_session()

    def export_settings_snapshot(self) -> None:
        app_data_dir()
        data = self.session_data()
        path, _ = QFileDialog.getSaveFileName(self, "Export settings snapshot", str(APP_DIR / f"settings_snapshot_{timestamp_slug()}.json"), "JSON (*.json)")
        if path:
            try:
                atomic_write_text(path, json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
                self.show_compact_message("Settings Snapshot", f"Settings exported:\n{path}", "info")
            except Exception as exc:
                self.show_compact_message("Settings Snapshot failed", str(exc), "error")

    def write_history(self, idx: int, result: JobResult) -> None:
        try:
            if not (0 <= idx < len(self.jobs)):
                return
            job = self.jobs[idx]
            rec = {
                "time": time.strftime("%Y-%m-%d %H:%M:%S"),
                "status": result.status,
                "service": job.service,
                "id": job.ident,
                "url": redact_sensitive_text(job.url),
                "rc": result.rc,
                "error_type": classify_error(result.message),
                "downloaded": result.downloaded,
                "skipped": result.skipped,
                "message": redact_sensitive_text(result.message),
            }
            self.feature_store.record_history(rec)
        except Exception as exc:
            reporter = getattr(self, "_report_persistence_failure", None)
            if callable(reporter):
                reporter("writing history", exc)

    def load_history_table(self) -> None:
        self.tbl_history.setRowCount(0)
        try:
            records = self.feature_store.list_history(limit=500)
        except Exception:
            return
        self.tbl_history.setRowCount(len(records))
        for row, rec in enumerate(records):
            self.tbl_history.setItem(row, 0, QTableWidgetItem(str(rec.get("occurred_at", ""))))
            self.tbl_history.setItem(row, 1, self.status_item(str(rec.get("status", ""))))
            self.tbl_history.setItem(row, 2, QTableWidgetItem(str(rec.get("service", ""))))
            self.tbl_history.setItem(row, 3, QTableWidgetItem(str(rec.get("item_id", ""))))
            self.tbl_history.setItem(row, 4, QTableWidgetItem(str(rec.get("url", ""))))

    def clear_history(self) -> None:
        res = QMessageBox.question(self, "Clear history", "Delete saved job history?")
        if res == QMessageBox.Yes:
            self.feature_store.clear_history()
            self.load_history_table()

    def notify_finished(self) -> None:
        if not getattr(self, "notifications_enabled", True):
            return
        if QSystemTrayIcon.isSystemTrayAvailable():
            if self.tray_icon is None:
                self.tray_icon = QSystemTrayIcon(self)
            # An icon is REQUIRED: without one, Windows logs "No Icon set" and
            # frequently refuses to show the balloon at all — the "finished"
            # notification silently never appeared. Fall back to a standard
            # style icon when the window has none.
            icon = self.windowIcon()
            if icon.isNull():
                icon = self.style().standardIcon(QStyle.SP_ComputerIcon)
            self.tray_icon.setIcon(icon)
            self.tray_icon.show()
            msg = "Antrean download selesai." if self._ui_is_indonesian() else "Download queue finished."
            self.tray_icon.showMessage(APP_NAME, msg, QSystemTrayIcon.Information, 3000)
            # Don't leave an orphan tray icon behind for the rest of the app's
            # lifetime once the balloon has expired.
            if not getattr(self, "close_to_tray", False):
                QTimer.singleShot(7000, lambda: self.tray_icon.hide() if self.tray_icon else None)

    def import_preview_dialog(self) -> None:
        if self.active_workers > 0:
            self.show_compact_message("Import Preview", "Stop the current download before previewing another database.", "warning")
            return
        path, _ = QFileDialog.getOpenFileName(
            self,
            "Preview link database",
            "",
            "Supported (*.txt *.csv *.xlsx);;Text (*.txt);;CSV (*.csv);;Excel (*.xlsx);;All files (*.*)",
        )
        if not path:
            return
        try:
            suffix = Path(path).suffix.lower()
            if suffix == ".xlsx":
                text = self.read_xlsx_as_commands(path)
            elif suffix == ".csv":
                text = self.read_csv_as_commands(path)
            else:
                text = read_text_safely(path)
            jobs = parse_text_database(text)
            services: dict[str, int] = {}
            for job in jobs:
                services[job.service] = services.get(job.service, 0) + 1
            sample = "\n".join(
                f"{i+1}. {redact_sensitive_text(j.raw)[:160]}"
                for i, j in enumerate(jobs[:8])
            ) or "-"
            service_lines = "\n".join(f"- {k}: {v}" for k, v in sorted(services.items(), key=lambda x: (-x[1], x[0]))[:12]) or "-"
            msg = (
                f"File: {path}\n"
                f"Parsed jobs: {len(jobs)}\n\n"
                f"Services:\n{service_lines}\n\n"
                f"Sample rows:\n{sample}\n\n"
                "This preview does not import or change your current queue."
            )
            self.show_scroll_message("Import Preview", msg, "info")
        except Exception as exc:
            self.show_compact_message("Import Preview Failed", str(exc), "error")

    def open_service_policy_editor(self) -> None:
        dlg = QDialog(self)
        dlg.setWindowTitle("Service Rate Policy")
        dlg.resize(620, 460)
        lay = QVBoxLayout(dlg)
        note = QLabel("Optional GUI scheduler policy. Format: service, worker limit, retry count, delay seconds. Leave service empty rows unused.")
        note.setWordWrap(True)
        lay.addWidget(note)
        # Size the table to the existing policy instead of a fixed 8 rows.
        # A session/profile can legitimately carry more than 8 services
        # (apply_session_data accepts any size); with the fixed table, opening
        # this editor and pressing Save silently DELETED every entry past row 8.
        existing = list(self.service_policy.items())
        row_count = max(8, len(existing) + 4)
        table = QTableWidget(row_count, 4)
        table.setHorizontalHeaderLabels(["Service", "Worker limit", "Retries", "Delay sec"])
        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        table.setColumnWidth(1, 95)
        table.setColumnWidth(2, 80)
        table.setColumnWidth(3, 90)
        for row, (service, policy) in enumerate(existing):
            table.setItem(row, 0, QTableWidgetItem(str(service)))
            table.setItem(row, 1, QTableWidgetItem(str(policy.get("workers", 0))))
            table.setItem(row, 2, QTableWidgetItem(str(policy.get("retries", self.spin_retries.value()))))
            table.setItem(row, 3, QTableWidgetItem(str(policy.get("delay", 0))))
        lay.addWidget(table, 1)
        btns = QHBoxLayout()
        btns.addStretch(1)
        save = QPushButton("Save Policy")
        close = QPushButton("Cancel")
        btns.addWidget(save)
        btns.addWidget(close)
        lay.addLayout(btns)
        close.clicked.connect(dlg.reject)
        def save_policy() -> None:
            policy: dict[str, dict[str, int]] = {}
            for r in range(table.rowCount()):
                service_item = table.item(r, 0)
                service = service_item.text().strip().lower() if service_item else ""
                if not service:
                    continue
                workers = safe_int(table.item(r, 1).text() if table.item(r, 1) else 0, 0, 0, 32)
                retries = safe_int(table.item(r, 2).text() if table.item(r, 2) else self.spin_retries.value(), self.spin_retries.value(), 0, 20)
                delay = safe_int(table.item(r, 3).text() if table.item(r, 3) else 0, 0, 0, 3600)
                policy[service] = {"workers": workers, "retries": retries, "delay": delay}
            self.service_policy = policy
            self.autosave_session()
            self.append_log(f"[policy] service policy saved: {self.service_policy}")
            dlg.accept()
        save.clicked.connect(save_policy)
        dlg.exec()

    def open_retry_strategy_editor(self) -> None:
        dlg = QDialog(self)
        dlg.setWindowTitle("Retry Strategy")
        dlg.resize(480, 360)
        lay = QVBoxLayout(dlg)
        note = QLabel("Choose which error types Retry Failed may retry. Auth, config, path, and not-found errors are usually unsafe to retry blindly.")
        note.setWordWrap(True)
        lay.addWidget(note)
        checks: dict[str, QCheckBox] = {}
        for name in ["network", "rate-limit", "unknown", "path", "config", "auth/cookies", "not-found"]:
            cb = QCheckBox(name)
            cb.setChecked(bool(self.retry_strategy.get(name, True)))
            checks[name] = cb
            lay.addWidget(cb)
        row = QHBoxLayout()
        row.addStretch(1)
        save = QPushButton("Save")
        cancel = QPushButton("Cancel")
        row.addWidget(save)
        row.addWidget(cancel)
        lay.addLayout(row)
        cancel.clicked.connect(dlg.reject)
        def save_strategy() -> None:
            self.retry_strategy = {k: cb.isChecked() for k, cb in checks.items()}
            self.autosave_session()
            self.append_log(f"[policy] retry strategy saved: {self.retry_strategy}")
            dlg.accept()
        save.clicked.connect(save_strategy)
        dlg.exec()

    def group_queue_dialog(self) -> None:
        if not self.jobs:
            self.show_compact_message("Group Queue", "No queue rows to analyze.", "info")
            return
        by_service: dict[str, int] = {}
        by_tag: dict[str, int] = {}
        for job in self.jobs:
            by_service[job.service or "-"] = by_service.get(job.service or "-", 0) + 1
            by_tag[job.tag or "untagged"] = by_tag.get(job.tag or "untagged", 0) + 1
        msg = "By service:\n" + "\n".join(f"- {k}: {v}" for k, v in sorted(by_service.items(), key=lambda x: (-x[1], x[0])))
        msg += "\n\nBy tag:\n" + "\n".join(f"- {k}: {v}" for k, v in sorted(by_tag.items(), key=lambda x: (-x[1], x[0])))
        msg += "\n\nUse the buttons below for actual sorting."
        dlg = QDialog(self)
        dlg.setWindowTitle("Group/Sort Queue")
        dlg.resize(560, 420)
        lay = QVBoxLayout(dlg)
        box = QPlainTextEdit()
        box.setReadOnly(True)
        box.setPlainText(msg)
        lay.addWidget(box, 1)
        row = QHBoxLayout()
        row.addStretch(1)
        by_s = QPushButton("Sort by Service")
        by_t = QPushButton("Sort by Tag")
        close = QPushButton("Close")
        row.addWidget(by_s)
        row.addWidget(by_t)
        row.addWidget(close)
        lay.addLayout(row)
        def sort_key_service(job: DownloadJob):
            return ((job.service or "-"), (job.tag or ""), job.url)
        def sort_key_tag(job: DownloadJob):
            return ((job.tag or ""), (job.service or "-"), job.url)
        def do_sort(keyfunc) -> None:
            if self.active_workers > 0:
                self.show_compact_message("Sort Queue", "Stop current download before sorting the queue.", "warning")
                return
            self.jobs.sort(key=keyfunc)
            self.sync_editor_from_jobs()
            self._clear_result_state()
            self.populate_queue_table()
            dlg.accept()
        by_s.clicked.connect(lambda: do_sort(sort_key_service))
        by_t.clicked.connect(lambda: do_sort(sort_key_tag))
        close.clicked.connect(dlg.accept)
        dlg.exec()

    def dry_run_validator(self) -> None:
        if not self._rebuild_from_text():
            return
        if not self.jobs:
            self.show_compact_message("Dry-run Validator", "No queue rows to validate.", "info")
            return
        # Pure command-builder use, mirroring final_command_for_job: private
        # events/cancel-set (never mutate live run state), Qt-owned lifetime.
        dummy_q: queue.Queue = queue.Queue()
        validator = DownloadWorker(
            0, dummy_q, self.gdl_cmd or "gallery-dl", self.config_path, self.edit_output.text().strip(),
            self.combo_cookies.currentText(), self.spin_retries.value(), "", threading.Event(), threading.Event(),
            set(), False, "zip", False, self.service_policy, {}, None, self,
        )
        errors: list[str] = []
        previews: list[str] = []
        try:
            for i, job in enumerate(self.jobs):
                try:
                    cmd = validator.build_command(job)
                    if not cmd:
                        errors.append(f"#{i+1}: empty command")
                    # Full gallery-dl commands can source URLs from -i/-I/-x
                    # or invoke extractor targets such as oauth:pixiv. They
                    # were already source-validated by parse_line(). Apply the
                    # HTTP shorthand warning only to non-command rows.
                    if (
                        not job.is_command
                        and not job.url.startswith(("http://", "https://", "oauth:"))
                    ):
                        errors.append(f"#{i+1}: no HTTP/HTTPS URL detected")
                    previews.append(f"#{i+1}: " + " ".join(quote_arg_for_preview(x) for x in redact_sensitive_argv(cmd)[:14]))
                except Exception as exc:
                    errors.append(f"#{i+1}: {exc}")
        finally:
            validator.deleteLater()
        msg = f"Checked rows: {len(self.jobs)}\nProblems: {len(errors)}\n\n"
        if errors:
            msg += "Problems:\n" + "\n".join(errors[:80]) + "\n\n"
        msg += "Preview:\n" + "\n".join(previews[:40])
        self.show_scroll_message("Dry-run Validator", msg, "warning" if errors else "info")

    def apply_output_folder_by_tag(self) -> None:
        if self.active_workers > 0:
            self.show_compact_message("Output Folder per Tag", "Stop current download before changing the input.", "warning")
            return
        if not self._rebuild_from_text():
            return
        if not self.jobs:
            self.show_compact_message("Output Folder per Tag", "No jobs found.", "info")
            return
        base = QFileDialog.getExistingDirectory(self, "Choose base output folder", self.edit_output.text().strip() or str(Path.home()))
        if not base:
            return
        changed = 0
        new_lines: list[str] = []
        # Re-emit tag headers while rewriting rows, INCLUDING a bare "#" reset
        # when an untagged job follows a tagged group — tags are sticky on
        # reparse, so omitting the reset silently tags untagged rows (same
        # emission rule as sync_editor_from_jobs).
        last_tag = ""
        for job in self.jobs:
            if job.tag != last_tag:
                new_lines.append(f"# {job.tag}" if job.tag else "#")
                last_tag = job.tag
            if job.notes:
                new_lines.append(f"#@notes {' '.join(job.notes.split())}")
            if job.dest != "-" or not job.tag.strip():
                new_lines.append(job.raw)
                continue
            dest = str(Path(base) / safe_filename(job.tag.strip()))
            new_lines.append(command_with_destination(job, dest))
            changed += 1
        if not changed:
            self.show_compact_message("Output Folder per Tag", "No tagged rows without destination were found.", "info")
            return
        res = QMessageBox.question(self, "Output Folder per Tag", f"Apply tag-based destination to {changed} row(s)?")
        if res != QMessageBox.Yes:
            return
        self.txt_commands.setPlainText("\n".join(new_lines))
        self.append_log(f"[queue] applied output folder per tag to {changed} row(s)")

    def archive_path_assistant(self) -> None:
        default = str(Path(self.edit_output.text().strip() or "./downloads").expanduser() / "gallery-dl-archive.sqlite3")
        path, _ = QFileDialog.getSaveFileName(self, "Choose gallery-dl archive database path", default, "SQLite DB (*.sqlite3 *.db);;All files (*.*)")
        if not path:
            return
        # JSON-escape the path for the config snippet. Building this outside the
        # f-string keeps the backslash replacement off the expression part, which
        # is a SyntaxError on Python 3.9-3.11 (only relaxed in 3.12 via PEP 701).
        json_path = json.dumps(path)
        snippet = (
            "Config snippet:\n"
            "{\n"
            "  \"extractor\": {\n"
            f"    \"archive\": {json_path}\n"
            "  }\n"
            "}\n\n"
            "Command option example:\n"
            f"gallery-dl --download-archive {quote_arg_for_preview(path)} <URL>\n\n"
            "Note: this is gallery-dl's own archive database. It prevents already-downloaded items from being downloaded again."
        )
        QApplication.clipboard().setText(snippet)
        self.show_scroll_message("Archive Path Assistant", snippet + "\n\nCopied to clipboard.", "info")

    def dragEnterEvent(self, event) -> None:
        if event.mimeData().hasUrls():
            event.acceptProposedAction()

    def dropEvent(self, event) -> None:
        if self.active_workers > 0:
            self.show_compact_message("Drop file", "Stop the current download before dropping a new database file.", "info")
            return
        urls = event.mimeData().urls()
        if not urls:
            return
        # Process EVERY dropped file — previously only urls[0] was read and the
        # rest were silently ignored. Databases are concatenated in drop order.
        # Plain-newline join on purpose: a "# filename" separator would be
        # parsed as a tag header and silently tag every job by file name.
        paths = [u.toLocalFile() for u in urls if u.toLocalFile()]
        if not paths:
            return
        texts: list[str] = []
        loaded_paths: list[str] = []
        errors: list[str] = []
        for path in paths:
            try:
                suffix = Path(path).suffix.lower()
                if suffix == ".xlsx":
                    text = self.read_xlsx_as_commands(path)
                elif suffix == ".csv":
                    text = self.read_csv_as_commands(path)
                else:
                    text = read_text_safely(path)
                if text.strip():
                    texts.append(text.strip())
                    loaded_paths.append(path)
            except Exception as exc:
                errors.append(f"{Path(path).name}: {exc}")
        if texts:
            # Join with a bare "#" tag RESET between files: tags are sticky on
            # reparse, so if file A ends inside a tagged group, file B's
            # leading untagged rows would silently inherit A's tag. A bare "#"
            # clears the tag without introducing a new one (a "# filename"
            # separator would tag every job by file name).
            combined_text = "\n#\n".join(texts)
            try:
                parse_text_database(combined_text)
            except ValueError as exc:
                self.show_compact_message("Drop failed", str(exc), "error")
                event.acceptProposedAction()
                return
            self.txt_commands.setPlainText(combined_text)
            self.current_file = loaded_paths[0]
            label = Path(loaded_paths[0]).name
            if len(loaded_paths) > 1:
                label += f" (+{len(loaded_paths) - 1} more)"
            self.lbl_loaded_file.setText(label)
            self.lbl_loaded_file.setToolTip("\n".join(loaded_paths))
            self.append_log(
                f"[input] dropped {len(loaded_paths)} file(s): "
                f"{', '.join(Path(p).name for p in loaded_paths)}"
            )
        if errors:
            self.show_compact_message("Drop failed", "\n".join(errors), "error")
        event.acceptProposedAction()

    def closeEvent(self, event) -> None:
        if getattr(self, "close_to_tray", False) and not getattr(self, "_force_quit", False):
            self._ensure_persistent_tray()
            self.hide()
            event.ignore()
            return
        # Ask before stopping GUI timers. Previously declining this prompt left
        # the still-running window with its batched log timer disabled, so new
        # worker output accumulated without appearing in the UI.
        if self.active_workers > 0:
            res = QMessageBox.question(self, "Close", "Downloads are still running. Stop and close?")
            if res != QMessageBox.Yes:
                event.ignore()
                return
        if self.active_workers > 0:
            self.stop_download()
            # First wait for a graceful exit. Any worker that has not finished is
            # forced down (kill the child process group, then wait again) so a
            # QThread is never destroyed while still running and no gallery-dl
            # child process is left orphaned after the window closes.
            for worker in self.workers:
                worker.wait(3000)
            for worker in self.workers:
                if worker.isRunning():
                    try:
                        worker.terminate_current_process()
                    except Exception:
                        pass
                    worker.wait(2000)
            still_running = [worker for worker in self.workers if worker.isRunning()]
            if still_running:
                QMessageBox.warning(
                    self,
                    "Close delayed",
                    "A background worker is still shutting down. The window was kept open "
                    "to avoid corrupting files or leaving a child process behind. Try closing again shortly.",
                )
                event.ignore()
                return
            # closeEvent blocks the GUI thread while waiting, so queued
            # job_done/finished_all signals may not run before the application
            # exits. Finalize any still-unprocessed rows and the run record here
            # to avoid offering an intentional Stop as crash recovery next time.
            closing_run_id = getattr(self, "current_run_id", None)
            if closing_run_id:
                unfinished = set(self.running_indices) - set(self._processed_in_run)
                for idx in unfinished:
                    try:
                        self.feature_store.mark_run_item(
                            closing_run_id,
                            idx,
                            "stopped",
                            return_code=-1,
                            message="stopped while closing application",
                        )
                    except Exception:
                        pass
                try:
                    self.feature_store.finish_run(closing_run_id, "stopped")
                except Exception:
                    pass
                self.current_run_id = None
        file_workers = [
            worker
            for worker in (
                getattr(self, "_backup_file_worker", None),
                getattr(self, "_scan_file_worker", None),
            )
            if worker is not None and worker.isRunning()
        ]
        for worker in file_workers:
            worker.stop()
        for worker in file_workers:
            worker.wait(5000)
        if any(worker.isRunning() for worker in file_workers):
            QMessageBox.warning(
                self,
                "Close delayed",
                "A filesystem operation is still stopping. The window was kept open "
                "so its output file remains valid. Try closing again shortly.",
            )
            event.ignore()
            return
        # Wait for any background probe (version check) so a parented QThread
        # is never destroyed while still running.
        try:
            if self.version_worker is not None and self.version_worker.isRunning():
                self.version_worker.stop()
                if not self.version_worker.wait(3000):
                    self.version_worker.force_stop()
                    self.version_worker.wait(2000)
        except Exception:
            pass
        if self.version_worker is not None and self.version_worker.isRunning():
            QMessageBox.warning(
                self,
                "Close delayed",
                "The gallery-dl probe is still shutting down. Try closing again shortly.",
            )
            event.ignore()
            return
        runtime_process = getattr(self, "_runtime_process", None)
        if runtime_process is not None and runtime_process.state() != QProcess.NotRunning:
            # Prevent the venv stage's finished callback from launching the pip
            # stage while the application is in the middle of shutting down.
            self._runtime_stage = ""
            runtime_process.terminate()
            if not runtime_process.waitForFinished(3000):
                runtime_process.kill()
                runtime_process.waitForFinished(2000)
        # Only stop UI timers after every background component has confirmed
        # shutdown. If shutdown had to be delayed above, the live window keeps
        # its debounce, log flush, ETA, and delayed-start behavior intact.
        for timer_name in (
            "delay_timer", "_text_debounce", "_log_flush_timer", "eta_timer",
        ):
            timer = getattr(self, timer_name, None)
            if timer is not None:
                try:
                    timer.stop()
                except Exception:
                    pass
        try:
            if getattr(self, "schedule_timer", None) is not None:
                self.schedule_timer.stop()
        except Exception:
            pass
        self.cleanup_run_account_config()
        if self.tray_icon is not None:
            self.tray_icon.hide()
        # closeEvent can block queued worker-finished signals while shutting
        # down. Restore a queue displaced by Scheduler here as a second, safe
        # lifecycle boundary before the final autosave.
        self._restore_queue_after_scheduled_run()
        self.autosave_session()
        event.accept()


__all__ = ['AdvancedToolsMixin']
