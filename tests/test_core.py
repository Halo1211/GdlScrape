import os
import queue
import tempfile
import subprocess
import sys
import threading
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import gallery_dl_app as gui
from PySide6.QtWidgets import QMessageBox
from gallery_dl_app.advanced_tools import AdvancedToolsMixin
from gallery_dl_app.composer import (
    ComposerState,
    _bounded_int,
    _read_json_config,
    apply_config_editor_drafts,
    build_composer_argv,
    build_composer_config,
    composer_oauth_target,
    config_option_definitions,
    config_defaults,
    parse_typed_config_value,
    pixiv_site_options,
    reddit_site_options,
    redact_auth_config,
    site_archive_options,
    validate_cookies_txt,
)
from gallery_dl_app.queue_controller import (
    MAX_XLSX_UNCOMPRESSED_BYTES,
    QueueControllerMixin,
    validate_xlsx_archive,
)
from gallery_dl_app.reports import (
    CancellableFileTask,
    ReportsMixin,
    is_restorable_autosave,
    scan_output_tree,
    write_app_data_backup,
)
from gallery_dl_app.secure_vault import SecretVault
from gallery_dl_app.system_tools import SystemToolsMixin, is_gallery_dl_version_output


def make_worker(**overrides):
    values = {
        "worker_id": 0,
        "task_queue": queue.Queue(),
        "gdl_cmd": "gallery-dl",
        "config_path": None,
        "output_dir": "D:/Downloads",
        "cookies_browser": "none",
        "retries": 3,
        "extra_args": "",
        "pause_event": threading.Event(),
        "stop_event": threading.Event(),
        "cancelled_indices": set(),
        "compress_enabled": False,
        "compress_format": "zip",
        "convert_png_webp": False,
    }
    values.update(overrides)
    return gui.DownloadWorker(**values)


class SecretVaultTests(unittest.TestCase):
    def test_corrupt_dpapi_vault_is_not_overwritten_when_setting_secret(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "secure_secrets.json"
            original = b'{"existing": "ciphertext"'
            path.write_bytes(original)
            vault = SecretVault(path)
            vault._keyring = None

            with (
                patch("gallery_dl_app.secure_vault.os.name", "nt"),
                patch.object(vault, "_dpapi_protect", return_value=b"new-ciphertext"),
            ):
                with self.assertRaisesRegex(RuntimeError, "preserved"):
                    vault.set("account:new", "secret")

            self.assertEqual(path.read_bytes(), original)

    def test_keyring_delete_failure_is_reported_to_caller(self):
        with tempfile.TemporaryDirectory() as folder:
            vault = SecretVault(Path(folder) / "secure_secrets.json")
            vault._keyring = MagicMock()
            vault._keyring.get_password.return_value = "stored-secret"
            vault._keyring.delete_password.side_effect = RuntimeError("keyring locked")

            with self.assertRaisesRegex(RuntimeError, "keyring locked"):
                vault.delete("account:example")

    def test_keyring_delete_is_idempotent_when_secret_is_absent(self):
        with tempfile.TemporaryDirectory() as folder:
            vault = SecretVault(Path(folder) / "secure_secrets.json")
            vault._keyring = MagicMock()
            vault._keyring.get_password.return_value = None

            vault.delete("account:missing")

            vault._keyring.delete_password.assert_not_called()


class ParserTests(unittest.TestCase):
    def test_utf32_text_database_is_decoded_without_nul_characters(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "utf32.txt"
            expected = "https://example.com/gallery/123\n"
            path.write_bytes(expected.encode("utf-32"))

            decoded = gui.read_text_safely(path)

        self.assertEqual(decoded, expected)
        self.assertNotIn("\x00", decoded)

    def test_false_boolean_enabled_cell_disables_import_row(self):
        self.assertIsNone(
            gui.build_raw_command_from_columns({
                "url": "https://example.com/a",
                "enabled": False,
            })
        )

    def test_non_finite_numeric_booleans_use_the_requested_default(self):
        for value in (float("nan"), float("inf"), float("-inf")):
            with self.subTest(value=value):
                self.assertFalse(gui.safe_bool(value, False))
                self.assertTrue(gui.safe_bool(value, True))

    def test_service_detection_uses_gallery_dl_category_for_domain_aliases(self):
        cases = {
            "https://x.com/artist/status/123456789": ("twitter", "123456789"),
            "https://mobile.twitter.com/artist/status/123456789": ("twitter", "123456789"),
            "https://old.reddit.com/r/pics/comments/abc123/title/": ("reddit", "-"),
        }
        for url, expected in cases.items():
            with self.subTest(url=url):
                self.assertEqual(gui.parse_url_meta(url), expected)

    def test_service_detection_tolerates_malformed_ipv6_url(self):
        self.assertEqual(gui.parse_url_meta("https://[invalid"), ("-", "-"))

    def test_service_detection_ignores_urls_embedded_in_query_values(self):
        url = (
            "https://example.com/gallery/111?next="
            "https://www.pixiv.net/users/222"
        )

        self.assertEqual(gui.parse_url_meta(url), ("example", "111"))

    def test_command_executable_availability_rejects_missing_path_and_directory(self):
        with tempfile.TemporaryDirectory() as folder:
            self.assertFalse(gui.command_executable_available(folder))
            self.assertFalse(gui.command_executable_available(str(Path(folder) / "missing.exe")))
        self.assertTrue(gui.command_executable_available(sys.executable))
        self.assertTrue(gui.command_executable_available(f'"{sys.executable}" -m gallery_dl'))

    def test_posix_command_path_must_have_execute_permission(self):
        with tempfile.TemporaryDirectory() as folder:
            command = Path(folder) / "gallery-dl"
            command.write_text("#!/bin/sh\n", encoding="utf-8")
            with (
                patch("gallery_dl_app.core.IS_WINDOWS", False),
                patch("gallery_dl_app.core.os.access", return_value=False),
            ):
                self.assertFalse(gui.command_executable_available(str(command)))

    def test_rate_limit_classification_wins_over_generic_connection_text(self):
        self.assertEqual(
            gui.classify_error("Connection failed: HTTP 429 Too Many Requests"),
            "rate-limit",
        )

    def test_specific_errors_win_over_generic_connection_wording(self):
        cases = {
            "Connection failed: HTTP 403 Forbidden": "auth/cookies",
            "Network response: HTTP 404 Not Found": "not-found",
            "Connection failed because config JSON is invalid": "config",
            "Connection failed: output path has no such file": "path",
        }
        for message, expected in cases.items():
            with self.subTest(message=message):
                self.assertEqual(gui.classify_error(message), expected)

    def test_rate_detection_does_not_match_inside_unrelated_words(self):
        self.assertEqual(
            gui.classify_error("Generated output has an unknown failure"),
            "unknown",
        )

    def test_windows_native_return_code_fits_qt_signed_int(self):
        self.assertEqual(gui.normalize_process_return_code(0xC000013A), -1073741510)
        self.assertEqual(gui.normalize_process_return_code(0), 0)

    def test_recognizes_supported_invocations(self):
        accepted = [
            "gallery-dl https://example.com/a",
            "gallery-dl.exe https://example.com/a",
            "python -m gallery_dl https://example.com/a",
            "python3.12 -m gallery-dl https://example.com/a",
            "py -m gallery_dl https://example.com/a",
        ]
        for command in accepted:
            with self.subTest(command=command):
                self.assertTrue(gui.parse_line(command).is_command)

    def test_queue_ignores_destination_only_gallery_dl_templates(self):
        text = (
            'gallery-dl -d "F:\\Rips\\Download\\artist" '
            'https://www.pixiv.net/en/users/123\n'
            'gallery-dl -d "F:\\Rips\\Download\\riria"   \n'
        )

        jobs = gui.parse_text_database(text)

        self.assertEqual(len(jobs), 1)
        self.assertEqual(jobs[0].url, "https://www.pixiv.net/en/users/123")

    def test_queue_keeps_gallery_dl_input_file_commands_without_inline_url(self):
        commands = (
            'gallery-dl -d "F:\\Rips\\Download\\batch" -i urls.txt',
            'gallery-dl --input-file=urls.txt',
            'gallery-dl -Icommented-urls.txt',
            'gallery-dl -x -',
            'gallery-dl oauth:pixiv',
        )

        for command in commands:
            with self.subTest(command=command):
                self.assertIsNotNone(gui.parse_line(command))

    def test_queue_ignores_input_file_option_without_a_value(self):
        self.assertIsNone(gui.parse_line("gallery-dl --input-file"))
        self.assertIsNone(gui.parse_line("gallery-dl --input-file="))

    def test_empty_attached_input_file_does_not_hide_later_source(self):
        commands = (
            "gallery-dl --input-file= https://target.example/gallery",
            "gallery-dl --input-file-comment= https://target.example/gallery",
            "gallery-dl --input-file-delete= https://target.example/gallery",
        )

        for command in commands:
            with self.subTest(command=command):
                job = gui.parse_line(command)
                self.assertIsNotNone(job)
                self.assertEqual(job.url, "https://target.example/gallery")

    def test_rejects_arbitrary_module_launcher(self):
        job = gui.parse_line("malware -m gallery_dl https://example.com/a")
        self.assertIsNotNone(job)
        self.assertFalse(job.is_command)
        self.assertEqual(job.url, "https://example.com/a")

    def test_shorthand_options_are_tokenized(self):
        job = gui.parse_line('https://example.com/a --range 1-10 -d "D:/Other Folder"')
        self.assertEqual(job.url, "https://example.com/a")
        self.assertEqual(job.dest, "D:/Other Folder")
        command = make_worker().build_command(job)
        self.assertEqual(
            command,
            [
                "gallery-dl",
                "--retries",
                "3",
                "https://example.com/a",
                "--range",
                "1-10",
                "-d",
                "D:/Other Folder",
            ],
        )

    def test_exact_destination_is_not_overridden_by_gui_default(self):
        job = gui.parse_line('gallery-dl -D "D:/Exact Folder" https://example.com/a')
        self.assertEqual(job.dest, "D:/Exact Folder")
        command = make_worker(output_dir="D:/GUI Default").build_command(job)
        self.assertIn("-D", command)
        self.assertNotIn("-d", command)
        self.assertNotIn("D:/GUI Default", command)

    def test_attached_short_destination_is_not_overridden(self):
        job = gui.parse_line("gallery-dl -dD:/Attached https://example.com/a")
        self.assertEqual(job.dest, "D:/Attached")
        command = make_worker(output_dir="D:/GUI Default").build_command(job)
        self.assertIn("-dD:/Attached", command)
        self.assertNotIn("D:/GUI Default", command)

    @unittest.skipUnless(os.name == "nt", "Windows destination normalization")
    def test_windows_destination_drops_unusable_trailing_space(self):
        job = gui.parse_line(
            'gallery-dl -d "F:\\Rips\\Download\\DokiDoMiki " '
            "https://example.com/a"
        )

        command = make_worker(output_dir="D:/GUI Default").build_command(job)

        destination_index = command.index("-d") + 1
        self.assertEqual(command[destination_index], "F:\\Rips\\Download\\DokiDoMiki")
        self.assertEqual(job.dest, "F:\\Rips\\Download\\DokiDoMiki")

    def test_destination_like_option_value_is_not_treated_as_output(self):
        job = gui.parse_line(
            'gallery-dl --exec "-d" https://example.com/a'
        )
        self.assertEqual(job.dest, "-")
        command = make_worker(output_dir="D:/GUI Default").build_command(job)
        self.assertIn("D:/GUI Default", command)

    def test_destination_normalization_does_not_reinterpret_other_option_values(self):
        commands = (
            "gallery-dl --exec -d https://example.com/gallery.",
            "gallery-dl --header -D https://example.com/gallery.",
        )

        for source in commands:
            with self.subTest(source=source):
                job = gui.parse_line(source)
                command = make_worker(output_dir="").build_command(job)

                self.assertIn("https://example.com/gallery.", command)

    def test_proxy_url_is_not_selected_as_job_url(self):
        job = gui.parse_line(
            "gallery-dl --proxy http://127.0.0.1:8080 https://example.com/target"
        )
        self.assertEqual(job.url, "https://example.com/target")

    def test_boolean_option_after_proxy_does_not_make_proxy_the_target(self):
        job = gui.parse_line(
            "gallery-dl --proxy http://127.0.0.1:8080 --verbose "
            "https://example.com/target"
        )
        self.assertEqual(job.url, "https://example.com/target")

    def test_option_value_url_is_not_selected_as_job_url(self):
        job = gui.parse_line(
            "gallery-dl -o proxy=https://proxy.example/path "
            "https://target.example/gallery/123"
        )
        self.assertEqual(job.url, "https://target.example/gallery/123")

    def test_urls_inside_exec_and_print_output_are_not_selected_as_targets(self):
        commands = [
            'gallery-dl --exec "notify https://callback.example/hook" '
            "https://target.example/gallery/123",
            "gallery-dl --print-to-file id https://logs.example/output "
            "https://target.example/gallery/123",
        ]
        for command in commands:
            with self.subTest(command=command):
                self.assertEqual(
                    gui.parse_line(command).url,
                    "https://target.example/gallery/123",
                )

    def test_new_gallery_dl_option_values_are_not_selected_as_targets(self):
        commands = (
            "gallery-dl -a https://agent.example/value "
            "https://target.example/gallery/123",
            "gallery-dl --cache-file https://cache.example/db "
            "https://target.example/gallery/123",
            'gallery-dl --config-json \'{"callback":"https://config.example/hook"}\' '
            "https://target.example/gallery/123",
            "gallery-dl --Print-to-file id https://logs.example/output "
            "https://target.example/gallery/123",
            "gallery-dl --file-filter https://filter.example/value "
            "https://target.example/gallery/123",
        )

        for command in commands:
            with self.subTest(command=command):
                self.assertEqual(
                    gui.parse_line(command).url,
                    "https://target.example/gallery/123",
                )

    def test_option_value_only_command_is_ignored_as_a_template(self):
        self.assertIsNone(
            gui.parse_line("gallery-dl --cache-file https://cache.example/db")
        )

    def test_variadic_list_extractor_values_are_not_download_sources(self):
        commands = (
            "gallery-dl --list-extractors pixiv "
            "https://target.example/not-a-download",
            "gallery-dl --list-extractors pixiv -- "
            "https://target.example/not-a-download",
            "gallery-dl -i urls.txt --list-extractors pixiv",
        )
        for command in commands:
            with self.subTest(command=command):
                self.assertIsNone(gui.parse_line(command))

    def test_url_parser_tracks_all_installed_options_that_consume_values(self):
        from gallery_dl import option

        missing: list[str] = []
        for action in option.build_parser()._actions:
            if not action.option_strings or action.nargs == 0:
                continue
            missing.extend(
                flag
                for flag in action.option_strings
                if flag not in gui.GALLERY_DL_OPTION_VALUE_COUNTS
            )

        self.assertEqual(missing, [])

        variadic = {
            flag
            for action in option.build_parser()._actions
            if action.nargs in {"*", "+"}
            for flag in action.option_strings
        }
        self.assertEqual(variadic, gui.GALLERY_DL_VARIADIC_VALUE_FLAGS)
        wrong_arity = []
        for action in option.build_parser()._actions:
            if (
                not action.option_strings
                or action.nargs == 0
                or action.nargs in {"*", "+"}
            ):
                continue
            expected = 1 if action.nargs is None else int(action.nargs)
            wrong_arity.extend(
                (flag, gui.GALLERY_DL_OPTION_VALUE_COUNTS[flag], expected)
                for flag in action.option_strings
                if gui.GALLERY_DL_OPTION_VALUE_COUNTS[flag] != expected
            )
        self.assertEqual(wrong_arity, [])

    def test_windows_reserved_tag_becomes_safe_folder_name(self):
        self.assertEqual(gui.safe_filename("CON"), "_CON")
        self.assertEqual(gui.safe_filename("lpt1.txt"), "_lpt1.txt")
        self.assertEqual(gui.safe_filename("normal tag"), "normal tag")

    def test_safe_text_reader_supports_utf16_and_rejects_binary(self):
        with tempfile.TemporaryDirectory() as folder:
            utf16 = Path(folder) / "urls.txt"
            utf16.write_text("https://example.com/a\n", encoding="utf-16")
            binary = Path(folder) / "wrong-file.bin"
            binary.write_bytes(b"not text\x00with nul")

            self.assertEqual(gui.read_text_safely(utf16).splitlines(), ["https://example.com/a"])
            with self.assertRaisesRegex(ValueError, "binary"):
                gui.read_text_safely(binary)

    def test_safe_text_reader_rejects_oversized_file_before_reading(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "huge.txt"
            with path.open("wb") as handle:
                handle.seek(gui.MAX_IMPORT_BYTES)
                handle.write(b"x")
            with self.assertRaisesRegex(ValueError, "too large"):
                gui.read_text_safely(path)

    def test_queue_parser_rejects_excess_jobs_and_oversized_rows(self):
        with patch("gallery_dl_app.core.MAX_QUEUE_JOBS", 2):
            with self.assertRaisesRegex(ValueError, "2 job"):
                gui.parse_text_database(
                    "https://example.com/a\nhttps://example.com/b\nhttps://example.com/c"
                )
        with patch("gallery_dl_app.core.MAX_COMMAND_LINE_CHARS", 8):
            with self.assertRaisesRegex(ValueError, "safety limit"):
                gui.parse_text_database("https://example.com/a")

    def test_queue_parser_limits_comment_only_source_rows(self):
        with patch("gallery_dl_app.core.MAX_QUEUE_SOURCE_ROWS", 2):
            with self.assertRaisesRegex(ValueError, "source exceeds"):
                gui.parse_text_database("# one\n# two\n# three\n")

    def test_atomic_write_preserves_existing_file_mode(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "config.json"
            path.write_text("old", encoding="utf-8")
            previous_mode = path.stat().st_mode & 0o7777
            with patch("gallery_dl_app.core.os.chmod") as chmod:
                gui.atomic_write_text(path, "new", encoding="utf-8")
            chmod.assert_called_once()
            self.assertEqual(chmod.call_args.args[1], previous_mode)
            self.assertEqual(path.read_text(encoding="utf-8"), "new")

    def test_atomic_writes_use_distinct_temporary_files_per_call(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "session.json"
            temporary_paths = []
            with patch(
                "gallery_dl_app.core.os.replace",
                side_effect=lambda source, _target: temporary_paths.append(Path(source)),
            ):
                gui.atomic_write_text(path, "first", encoding="utf-8")
                gui.atomic_write_text(path, "second", encoding="utf-8")

        self.assertEqual(len(temporary_paths), 2)
        self.assertNotEqual(temporary_paths[0], temporary_paths[1])


class CommandBuilderTests(unittest.TestCase):
    def test_gui_defaults_do_not_duplicate_explicit_options(self):
        job = gui.parse_line(
            'gallery-dl -c "D:/custom.json" -d "D:/custom" -R 9 '
            "--cookies-from-browser edge https://example.com/a"
        )
        command = make_worker(
            config_path="D:/ignored.json",
            cookies_browser="firefox",
        ).build_command(job)
        self.assertEqual(command.count("-c"), 1)
        self.assertEqual(command.count("-d"), 1)
        self.assertEqual(command.count("-R"), 1)
        self.assertEqual(command.count("--cookies-from-browser"), 1)
        self.assertNotIn("--retries", command)

    def test_attached_short_options_override_gui_defaults(self):
        with tempfile.TemporaryDirectory() as folder:
            gui_config = Path(folder) / "gui.json"
            gui_config.write_text("{}", encoding="utf-8")
            job = gui.parse_line(
                "gallery-dl -cD:/custom.json -R5 -Ccookies.txt https://example.com/a"
            )
            command = make_worker(
                config_path=str(gui_config),
                cookies_browser="firefox",
                retries=9,
            ).build_command(job)
        self.assertNotIn("--config", command)
        self.assertNotIn("--cookies-from-browser", command)
        self.assertNotIn("--retries", command)

    def test_config_ignore_prevents_gui_config_injection(self):
        with tempfile.TemporaryDirectory() as folder:
            gui_config = Path(folder) / "gui.json"
            gui_config.write_text("{}", encoding="utf-8")
            job = gui.parse_line("gallery-dl --config-ignore https://example.com/a")
            command = make_worker(config_path=str(gui_config)).build_command(job)
        self.assertEqual(command.count("--config-ignore"), 1)
        self.assertNotIn("--config", command)

    def test_preview_redacts_secrets_without_mutating_source(self):
        source = ["gallery-dl", "--password", "secret", "--option", "api-key=123"]
        redacted = gui.redact_sensitive_argv(source)
        self.assertEqual(source[2], "secret")
        self.assertEqual(redacted[2], gui.REDACTED)
        self.assertEqual(redacted[4], f"api-key={gui.REDACTED}")

    def test_preview_redacts_attached_options_and_signed_urls(self):
        source = [
            "gallery-dl",
            "-psecret",
            "-uaccount",
            "-Ccookies.txt",
            "--option=extractor.api-key=abc123",
            "-orefresh-token=token123",
            "https://example.com/file?access_token=url-secret&item=42",
        ]
        redacted = gui.redact_sensitive_argv(source)
        self.assertEqual(source[1], "-psecret")
        self.assertEqual(redacted[1], f"-p{gui.REDACTED}")
        self.assertEqual(redacted[2], f"-u{gui.REDACTED}")
        self.assertEqual(redacted[3], f"-C{gui.REDACTED}")
        self.assertEqual(redacted[4], f"--option=extractor.api-key={gui.REDACTED}")
        self.assertEqual(redacted[5], f"-orefresh-token={gui.REDACTED}")
        self.assertEqual(
            redacted[6],
            f"https://example.com/file?access_token={gui.REDACTED}&item=42",
        )

    def test_preview_redacts_secrets_nested_inside_option_values(self):
        argv = [
            "gallery-dl",
            "--config-json",
            '{"password":"json-secret"}',
            "--exec",
            "notify --password callback-secret",
            "https://example.com",
        ]

        rendered = " ".join(gui.redact_sensitive_argv(argv))

        self.assertNotIn("json-secret", rendered)
        self.assertNotIn("callback-secret", rendered)
        self.assertIn(gui.REDACTED, rendered)

        attached_long = gui.redact_sensitive_argv([
            "gallery-dl",
            '--option=extractor.safe={"password":"nested-long-secret"}',
        ])[1]
        attached_short = gui.redact_sensitive_argv([
            "gallery-dl",
            '-oextractor.safe={"password":"nested-short-secret"}',
        ])[1]
        self.assertNotIn("nested-long-secret", attached_long)
        self.assertNotIn("nested-short-secret", attached_short)
        self.assertTrue(attached_long.startswith("--option="))
        self.assertTrue(attached_short.startswith("-oextractor.safe="))

    def test_structured_secret_redaction_handles_escaped_quotes(self):
        samples = (
            r'{"password":"prefix\"double-tail-secret"}',
            r"{'password':'prefix\'single-tail-secret'}",
            r'--config-json={"password":"prefix\"config-tail-secret"}',
        )

        for source in samples:
            with self.subTest(source=source):
                redacted = gui.redact_sensitive_text(source)
                self.assertNotIn("tail-secret", redacted)
                self.assertIn(gui.REDACTED, redacted)

    def test_signed_url_credentials_are_redacted_without_false_positive_keys(self):
        signed_urls = [
            (
                "https://bucket.s3.amazonaws.com/a?X-Amz-Algorithm=AWS4-HMAC-SHA256&"
                "X-Amz-Credential=AKIA_TEST/path&X-Amz-Signature=aws-secret",
                ("AKIA_TEST", "aws-secret"),
            ),
            (
                "https://storage.googleapis.com/a?X-Goog-Credential=user@example.com/path&"
                "X-Goog-Signature=google-secret",
                ("user@example.com", "google-secret"),
            ),
            (
                "https://x.blob.core.windows.net/a?sv=2024&sig=azure-secret",
                ("azure-secret",),
            ),
            (
                "https://cdn.example/a?Expires=123&Signature=cloudfront-secret",
                ("cloudfront-secret",),
            ),
            (
                "https://bucket.example/a?X%2DAmz%2DSignature=encoded-secret",
                ("encoded-secret",),
            ),
            (
                "https://example.com/oauth#access_token=fragment-secret&state=keep",
                ("fragment-secret",),
            ),
        ]
        for source, secrets in signed_urls:
            with self.subTest(source=source):
                redacted = gui.redact_sensitive_text(source)
                for secret in secrets:
                    self.assertNotIn(secret, redacted)
                self.assertIn(gui.REDACTED, redacted)

        innocent = (
            "https://example.com/search?tokenizer=wordpiece&"
            "username_display=true&item=42"
        )
        self.assertEqual(gui.redact_sensitive_text(innocent), innocent)

    def test_structured_log_secrets_are_redacted_without_corrupting_safe_options(self):
        sources = [
            "config {'password': 'hunter2', 'ok': 1}",
            '{"api-key": "abc123", "status": "ok"}',
            "extractor.password = hunter2",
            "password=hunter2 username=alice access_token=abc",
            "login token: abcdef",
        ]
        for source in sources:
            with self.subTest(source=source):
                redacted = gui.redact_sensitive_text(source)
                self.assertIn(gui.REDACTED, redacted)
                for secret in ("hunter2", "abc123", "alice", "abcdef"):
                    self.assertNotIn(secret, redacted)

        safe = (
            "gallery-dl -o tokenizer=wordpiece -o username_display=true "
            "-o parent-session=true https://example.com/a"
        )
        self.assertEqual(gui.redact_sensitive_text(safe), safe)
        self.assertFalse(gui.is_sensitive_option_key("username_display"))
        self.assertFalse(gui.is_sensitive_option_key("parent-session"))

        signed_option = (
            "gallery-dl -o x-amz-signature=cloud-option-secret "
            "https://example.com/a"
        )
        redacted_option = gui.redact_sensitive_text(signed_option)
        self.assertNotIn("cloud-option-secret", redacted_option)
        self.assertTrue(redacted_option.endswith("https://example.com/a"))

    def test_redaction_hides_url_userinfo_and_secrets_inside_log_text(self):
        source = "gallery-dl -psecret --proxy=http://user:pass@proxy.test/?token=abc"
        redacted = gui.redact_sensitive_text(source)
        self.assertNotIn("secret", redacted)
        self.assertNotIn("user:pass", redacted)
        self.assertNotIn("token=abc", redacted)
        self.assertIn(gui.REDACTED, redacted)

    def test_redaction_hides_socks_proxy_userinfo(self):
        source = "--proxy=socks5://alice:secret@proxy.test:1080"
        redacted = gui.redact_sensitive_text(source)
        self.assertNotIn("alice:secret", redacted)
        self.assertIn(gui.REDACTED, redacted)

    def test_redaction_hides_complete_quoted_secret_values(self):
        sources = [
            'gallery-dl --password "top secret value" https://example.com/gallery/42',
            'gallery-dl -p "top secret value" https://example.com/gallery/42',
            'gallery-dl --option password="top secret value" https://example.com/gallery/42',
            'gallery-dl --option "password=top secret value" https://example.com/gallery/42',
        ]
        for source in sources:
            with self.subTest(source=source):
                redacted = gui.redact_sensitive_text(source)
                self.assertNotIn("top secret value", redacted)
                self.assertNotIn("secret value", redacted)
                self.assertIn(gui.REDACTED, redacted)
                self.assertTrue(redacted.endswith("https://example.com/gallery/42"))
        quoted_pair = gui.redact_sensitive_text(sources[-1])
        self.assertIn(f'"password={gui.REDACTED}"', quoted_pair)

    def test_redaction_fails_closed_for_unterminated_quoted_secret(self):
        redacted = gui.redact_sensitive_text(
            'gallery-dl --password "top secret value still being typed'
        )
        self.assertNotIn("top secret", redacted)
        self.assertNotIn("still being typed", redacted)
        self.assertTrue(redacted.endswith(gui.REDACTED))

    def test_redaction_hides_secrets_in_python_argv_representations(self):
        sources = [
            "Command ['gallery-dl', '--password', 'list-secret'] timed out",
            'Command ["gallery-dl", "--api-key=assigned-secret"] failed',
            "Command ['gallery-dl', '-pattached-secret'] failed",
            "Command ['gallery-dl', '-o', 'oauth-token=option-secret'] failed",
        ]
        for source in sources:
            with self.subTest(source=source):
                redacted = gui.redact_sensitive_text(source)
                self.assertNotIn("secret", redacted.replace(gui.REDACTED, ""))
                self.assertIn(gui.REDACTED, redacted)

    def test_database_redaction_preserves_metadata_and_hides_secrets(self):
        source = (
            "# Project\n#@notes keep\n"
            "gallery-dl --password secret https://example.com/a?token=url-secret"
        )
        redacted = gui.redact_sensitive_database_text(source)
        self.assertTrue(redacted.startswith("# Project\n#@notes keep\n"))
        self.assertNotIn(" secret", redacted)
        self.assertNotIn("url-secret", redacted)
        self.assertIn(gui.REDACTED, redacted)

    def test_quoted_authorization_header_redaction_preserves_target_argv(self):
        sources = [
            'gallery-dl --header "Authorization: Bearer header-secret" '
            "https://example.com/gallery/42",
            "gallery-dl --header Authorization:header-secret "
            "https://example.com/gallery/42",
            "gallery-dl --header=Authorization:header-secret "
            "https://example.com/gallery/42",
        ]
        for source in sources:
            with self.subTest(source=source):
                redacted = gui.redact_sensitive_database_text(source)
                self.assertNotIn("header-secret", redacted)
                self.assertIn("https://example.com/gallery/42", redacted)
                self.assertEqual(
                    gui.split_command(redacted)[-1],
                    "https://example.com/gallery/42",
                )

    def test_header_and_option_redaction_preserve_adjacent_fields(self):
        structured = (
            "request {'Authorization': 'Bearer dict-secret', 'status': 'ok'}"
        )
        option = (
            "gallery-dl -o headers.Cookie=session-secret "
            "https://example.com/a"
        )

        redacted_structured = gui.redact_sensitive_text(structured)
        redacted_option = gui.redact_sensitive_text(option)

        self.assertNotIn("dict-secret", redacted_structured)
        self.assertIn("'status': 'ok'", redacted_structured)
        self.assertNotIn("session-secret", redacted_option)
        self.assertIn("https://example.com/a", redacted_option)

    def test_apply_option_to_database_preserves_metadata_rows(self):
        source = "# Project\n#@notes keep this\nhttps://example.com/a\n\nhttps://example.com/b"
        updated = gui.append_extra_args_to_database_text(source, "--simulate")
        self.assertIn("# Project\n#@notes keep this\n", updated)
        self.assertNotIn("# Project --simulate", updated)
        self.assertEqual(updated.count("--simulate"), 2)

    def test_tag_destination_rewrite_preserves_shorthand_options(self):
        job = gui.parse_line("https://example.com/a --range 2-8 --simulate")
        command = gui.command_with_destination(job, "D:/Tagged Output")
        argv = gui.split_command(command)
        self.assertEqual(argv[:3], ["gallery-dl", "-d", "D:/Tagged Output"])
        self.assertIn("--range", argv)
        self.assertIn("2-8", argv)
        self.assertIn("--simulate", argv)

    def test_spreadsheet_formula_triggers_are_neutralized(self):
        for value in ("=CMD()", "+1+1", "-2+3", "@SUM(A1:A2)", "\t=1", "\r=1"):
            with self.subTest(value=value):
                self.assertEqual(gui.sanitize_spreadsheet_cell(value), "'" + value)
        self.assertEqual(gui.sanitize_spreadsheet_cell("https://example.com"), "https://example.com")

    def test_formula_escape_round_trip_does_not_leave_apostrophe_argument(self):
        for value in ("\t--simulate", "\r--simulate"):
            with self.subTest(value=value):
                command = gui.build_raw_command_from_columns({
                    "url": "https://example.com/a",
                    "extra_args": gui.sanitize_spreadsheet_cell(value),
                })
                argv = gui.split_command(command)
                self.assertIn("--simulate", argv)
                self.assertNotIn("'", argv)


class DownloadComposerTests(unittest.TestCase):
    def test_oauth_target_accepts_only_picker_sites_and_requires_mastodon_instance(self):
        self.assertEqual(composer_oauth_target("pixiv"), "oauth:pixiv")
        self.assertEqual(
            composer_oauth_target("mastodon", "https://mastodon.social/"),
            "oauth:mastodon:https://mastodon.social/",
        )
        with self.assertRaisesRegex(ValueError, "supported OAuth site"):
            composer_oauth_target("pixvi")
        with self.assertRaisesRegex(ValueError, "instance"):
            composer_oauth_target("mastodon")
        with self.assertRaisesRegex(ValueError, "hostname or URL"):
            composer_oauth_target("mastodon", "https://user:pass@mastodon.social/callback?x=1")

    def test_per_site_catalog_combines_general_and_extractor_specific_options(self):
        pixiv = {item.key: item for item in config_option_definitions("pixiv")}
        kemono = {item.key: item for item in config_option_definitions("kemono")}
        reddit = {item.key: item for item in config_option_definitions("reddit")}

        self.assertEqual(pixiv["archive"].scope, "general")
        self.assertEqual(pixiv["ugoira"].scope, "site")
        self.assertEqual(pixiv["include"].value_type, "json")
        self.assertIn("duplicates", kemono)
        self.assertIn("archives-format", kemono)
        self.assertIn("client-id", reddit)
        self.assertTrue(reddit["client-id"].sensitive)
        self.assertGreater(len(config_option_definitions("pixiv")), 30)
        twitter = {item.key: item for item in config_option_definitions("twitter")}
        self.assertEqual((twitter["replies"].value_type, twitter["replies"].default), ("boolean", True))

    def test_config_editor_applies_general_and_site_edits_and_real_removals(self):
        original = {
            "extractor": {
                "retries": 4,
                "unknown-global": "keep",
                "pixiv": {
                    "metadata": False,
                    "refresh-token": "remove-me",
                    "unknown-site": "keep",
                },
            },
            "downloader": {"part": True},
        }
        result = apply_config_editor_drafts(
            original,
            general_overrides={"timeout": 45},
            site_overrides={"pixiv": {"metadata": True}},
            general_removals={"retries"},
            site_removals={"pixiv": {"refresh-token"}},
        )

        self.assertEqual(result["extractor"]["timeout"], 45)
        self.assertNotIn("retries", result["extractor"])
        self.assertIs(result["extractor"]["pixiv"]["metadata"], True)
        self.assertNotIn("refresh-token", result["extractor"]["pixiv"])
        self.assertEqual(result["extractor"]["unknown-global"], "keep")
        self.assertEqual(result["extractor"]["pixiv"]["unknown-site"], "keep")
        self.assertEqual(result["downloader"], {"part": True})

    def test_site_archive_gui_builds_current_gallery_dl_keys(self):
        options = site_archive_options(
            "E:/RIPS/Database/kemono.sqlite3",
            duplicates=False,
            archive_format=r"{service}{user}{id}\_{num}",
        )

        self.assertEqual(options["archive"], "E:/RIPS/Database/kemono.sqlite3")
        self.assertIs(options["duplicates"], False)
        self.assertEqual(options["archive-format"], "{service}{user}{id}_{num}")

    def test_reddit_gui_uses_oauth_specific_user_agent_key(self):
        options = reddit_site_options("client-id-value", "Python:Downloader:v1.01 (by /u/user)")

        self.assertEqual(options["client-id"], "client-id-value")
        self.assertEqual(options["user-agent-oauth"], "Python:Downloader:v1.01 (by /u/user)")
        self.assertNotIn("user-agent", options)

    def test_pixiv_gui_builds_requested_non_secret_options(self):
        options = pixiv_site_options(
            include=["background", "artworks", "avatar", "novel-user"],
            embeds=True,
            covers=True,
            full_series=True,
            metadata=True,
            ugoira="original",
        )

        self.assertEqual(options["include"], ["background", "artworks", "avatar", "novel-user"])
        self.assertIs(options["embeds"], True)
        self.assertEqual(options["ugoira"], "original")

    def test_advanced_site_option_parser_is_typed_and_rejects_invalid_json(self):
        self.assertIs(parse_typed_config_value("false", "boolean"), False)
        self.assertEqual(parse_typed_config_value("42", "integer"), 42)
        self.assertEqual(parse_typed_config_value('["a", "b"]', "json"), ["a", "b"])
        self.assertIsNone(parse_typed_config_value("", "null"))
        with self.assertRaisesRegex(ValueError, "valid JSON"):
            parse_typed_config_value("[broken", "json")

    def test_bounded_config_integer_handles_overflowing_json_number(self):
        self.assertEqual(_bounded_int(1e309, 0, 99), 0)
        self.assertEqual(_bounded_int(-1e309, 5, 99), 5)

    def test_json_config_reader_accepts_utf8_bom(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "config.json"
            path.write_text('{"extractor": {}}', encoding="utf-8-sig")
            data, error = _read_json_config(str(path))

        self.assertIsNone(error)
        self.assertEqual(data, {"extractor": {}})

    def test_json_config_reader_rejects_non_object_root(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "config.json"
            path.write_text("[]", encoding="utf-8")
            data, error = _read_json_config(str(path))

        self.assertEqual(data, {})
        self.assertIn("object", error)

    def test_json_config_reader_rejects_directory_path(self):
        with tempfile.TemporaryDirectory() as folder:
            data, error = _read_json_config(folder)

        self.assertEqual(data, {})
        self.assertIn("not a file", error)

    def test_open_destination_rejects_existing_file(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "not-a-folder.txt"
            target.write_text("data", encoding="utf-8")
            harness = MagicMock()
            harness.selected_indices.return_value = [0]
            harness.jobs = [MagicMock(dest=str(target))]

            with patch("gallery_dl_app.queue_controller.open_path") as opener:
                QueueControllerMixin.open_selected_destination(harness)

        opener.assert_not_called()
        harness.show_compact_message.assert_called_once()
        self.assertEqual(harness.show_compact_message.call_args.args[-1], "warning")

    def test_job_scope_emits_shared_values_as_overrides(self):
        state = ComposerState(
            destination="D:/Media",
            directory=["{category}", "{user[id]}"],
            filename="{id}.{extension}",
            cookies_browser="firefox",
            archive_enabled=True,
            archive_path="D:/archive.sqlite3",
            retries=5,
            timeout=30,
            file_range="1-20",
        )
        argv = build_composer_argv(state, "https://example.com/a")
        self.assertIn("-d", argv)
        self.assertIn("-f", argv)
        self.assertIn("--cookies-from-browser", argv)
        self.assertIn("--download-archive", argv)
        self.assertIn("--range", argv)

    def test_compatibility_overrides_use_stable_config_option_form(self):
        state = ComposerState(
            exact_destination=True,
            destination="D:/Exact",
            directory=["ignored", "for exact destination"],
            sleep_429="60-180",
            date_after="2025-01-01",
            date_before="2025-12-31",
            user_agent="browser",
            windows_filenames=True,
        )
        argv = build_composer_argv(state, "https://example.com/a")
        self.assertIn("-D", argv)
        self.assertNotIn('--windows-filenames', argv)
        self.assertNotIn('--sleep-429', argv)
        self.assertNotIn('--date-after', argv)
        self.assertNotIn('--date-before', argv)
        self.assertNotIn('-a', argv)
        self.assertIn('--user-agent', argv)
        self.assertIn("sleep-429=60-180", argv)
        self.assertIn("date-min=2025-01-01", argv)
        self.assertIn("date-max=2025-12-31", argv)
        self.assertIn("path-restrict=windows", argv)
        self.assertFalse(any(value.startswith("directory=") for value in argv))

    def test_config_scope_uses_saved_defaults_without_duplicate_cli_values(self):
        with tempfile.NamedTemporaryFile(suffix=".json") as config:
            state = ComposerState(
                apply_to_config=True,
                destination="D:/Media",
                filename="{id}.{extension}",
                cookies_browser="firefox",
                file_range="1-20",
            )
            argv = build_composer_argv(state, "https://example.com/a", config_path=config.name)
        self.assertIn("--config", argv)
        self.assertNotIn("-d", argv)
        self.assertNotIn("-f", argv)
        self.assertNotIn("--cookies-from-browser", argv)
        self.assertIn("--range", argv)

    def test_composer_does_not_pass_a_directory_as_config_file(self):
        with tempfile.TemporaryDirectory() as folder:
            state = ComposerState(use_active_config=True)
            argv = build_composer_argv(
                state,
                "https://example.com/a",
                config_path=folder,
            )

        self.assertNotIn("--config", argv)

    def test_config_merge_preserves_unknown_options(self):
        existing = {
            "extractor": {
                "custom-option": 42,
                "pixiv": {"refresh-token": "keep"},
                "postprocessors": [{"name": "exec", "command": "notify"}],
            },
            "output": {"progress": False},
        }
        state = ComposerState(destination="D:/Media", retries=4, metadata_json=True, archive_format="zip")
        merged = build_composer_config(
            state,
            site_blocks={"pixiv": {"filename": "{id}.{extension}"}},
            existing=existing,
        )
        self.assertEqual(merged["extractor"]["custom-option"], 42)
        self.assertEqual(merged["extractor"]["pixiv"]["refresh-token"], "keep")
        self.assertEqual(merged["extractor"]["pixiv"]["filename"], "{id}.{extension}")
        self.assertFalse(merged["output"]["progress"])
        self.assertEqual(
            [item["name"] for item in merged["extractor"]["postprocessors"]],
            ["exec", "metadata", "zip"],
        )

    def test_existing_config_defaults_are_read_back(self):
        values = config_defaults({
            "extractor": {
                "base-directory": "D:/Media",
                "directory": ["{category}", "{id}"],
                "cookies": ["edge"],
                "archive": "D:/archive.sqlite3",
                "postprocessors": [
                    {"name": "metadata", "event": "init", "filename": "info.json"},
                    {"name": "zip", "extension": "cbz"},
                ],
            }
        })
        self.assertEqual(values["destination"], "D:/Media")
        self.assertEqual(values["directory"], ["{category}", "{id}"])
        self.assertEqual(values["cookies_browser"], "edge")
        self.assertTrue(values["archive_enabled"])
        self.assertTrue(values["info_json"])
        self.assertFalse(values["metadata_json"])
        self.assertEqual(values["archive_format"], "cbz")

    def test_exact_destination_config_is_read_back(self):
        values = config_defaults({
            "extractor": {
                "base-directory": "D:/Exact",
                "directory": [],
            }
        })
        self.assertTrue(values["exact_destination"])

    def test_cleared_composer_settings_are_removed_but_unknowns_survive(self):
        existing = {
            "extractor": {
                "base-directory": "D:/Old",
                "directory": ["old"],
                "filename": "old.{extension}",
                "archive": "D:/old.sqlite3",
                "retries": 9,
                "cookies": ["firefox"],
                "custom-option": 42,
                "pixiv": {"refresh-token": "keep"},
                "postprocessors": [
                    {"name": "exec", "command": "notify"},
                    {"name": "metadata"},
                    {"name": "zip"},
                ],
            },
            "output": {"progress": False},
        }
        merged = build_composer_config(ComposerState(), existing=existing)
        extractor = merged["extractor"]
        for key in ("base-directory", "directory", "filename", "archive", "retries"):
            self.assertNotIn(key, extractor)
        self.assertEqual(extractor["cookies"], ["firefox"])
        self.assertEqual(extractor["custom-option"], 42)
        self.assertEqual(extractor["pixiv"]["refresh-token"], "keep")
        self.assertEqual(extractor["postprocessors"], [{"name": "exec", "command": "notify"}])
        self.assertFalse(merged["output"]["progress"])

    def test_site_credentials_are_config_only_and_redacted_in_preview(self):
        state = ComposerState(
            auth_category="danbooru",
            username="account-name",
            secret_key="password",
            secret_value="api-key-value",
            extra_auth_key="client-secret",
            extra_auth_value="second-secret",
        )
        config = build_composer_config(state)
        self.assertEqual(config["extractor"]["danbooru"]["username"], "account-name")
        self.assertEqual(config["extractor"]["danbooru"]["password"], "api-key-value")
        argv = build_composer_argv(state, "https://example.com/a")
        self.assertNotIn("api-key-value", argv)
        safe = redact_auth_config(config, extra_keys=(state.secret_key, state.extra_auth_key))
        self.assertEqual(safe["extractor"]["danbooru"]["username"], "<hidden>")
        self.assertEqual(safe["extractor"]["danbooru"]["password"], "<hidden>")
        self.assertEqual(safe["extractor"]["danbooru"]["client-secret"], "<hidden>")

    def test_switching_saved_auth_method_removes_stale_managed_credentials(self):
        existing = {
            "extractor": {
                "pixiv": {
                    "cookies": ["firefox"],
                    "cookies-update": True,
                    "username": "old-user",
                    "password": "old-secret",
                    "oauth-token": "old-custom-token",
                    "parent-session": True,
                    "username-display": "public-label",
                    "custom-option": "preserved",
                }
            }
        }
        credential_state = ComposerState(
            apply_to_config=True,
            auth_category="pixiv",
            username="new-user",
            secret_key="password",
            secret_value="new-secret",
        )

        credential_config = build_composer_config(
            credential_state,
            existing=existing,
        )
        pixiv = credential_config["extractor"]["pixiv"]
        self.assertNotIn("cookies", pixiv)
        self.assertNotIn("cookies-update", pixiv)
        self.assertNotIn("oauth-token", pixiv)
        self.assertEqual(pixiv["username"], "new-user")
        self.assertEqual(pixiv["password"], "new-secret")
        self.assertIs(pixiv["parent-session"], True)
        self.assertEqual(pixiv["username-display"], "public-label")
        self.assertEqual(pixiv["custom-option"], "preserved")

        browser_state = ComposerState(
            apply_to_config=True,
            auth_category="pixiv",
            cookies_browser="edge",
        )
        browser_config = build_composer_config(
            browser_state,
            existing=credential_config,
        )
        pixiv = browser_config["extractor"]["pixiv"]
        self.assertEqual(pixiv["cookies"], ["edge"])
        self.assertNotIn("username", pixiv)
        self.assertNotIn("password", pixiv)
        self.assertIs(pixiv["parent-session"], True)
        self.assertEqual(pixiv["username-display"], "public-label")
        self.assertEqual(pixiv["custom-option"], "preserved")

    def test_sensitive_key_detection_does_not_match_unrelated_substrings(self):
        self.assertTrue(gui.is_sensitive_option_key("oauth-token"))
        self.assertTrue(gui.is_sensitive_option_key("service.api_key"))
        self.assertFalse(gui.is_sensitive_option_key("tokenizer"))
        self.assertFalse(gui.is_sensitive_option_key("username_display"))
        self.assertFalse(gui.is_sensitive_option_key("parent-session"))
        self.assertFalse(gui.is_sensitive_option_key("custom-option"))

        preview = redact_auth_config({
            "oauth-token": "secret",
            "tokenizer": "wordpiece",
            "username_display": True,
            "parent-session": True,
        })
        self.assertEqual(preview["oauth-token"], "<hidden>")
        self.assertEqual(preview["tokenizer"], "wordpiece")
        self.assertIs(preview["username_display"], True)
        self.assertIs(preview["parent-session"], True)

    def test_browser_profile_and_domain_are_scoped_to_site(self):
        state = ComposerState(
            auth_category="twitter",
            cookies_browser="firefox",
            cookies_profile="Personal",
            cookies_domain=".twitter.com",
        )
        config = build_composer_config(state)
        self.assertEqual(
            config["extractor"]["twitter"]["cookies"],
            ["firefox", "Personal", None, None, ".twitter.com"],
        )
        argv = build_composer_argv(state, "https://twitter.com/example")
        self.assertIn("firefox/.twitter.com:Personal", argv)

    def test_cookies_txt_validator_accepts_netscape_and_rejects_json(self):
        with tempfile.TemporaryDirectory() as folder:
            valid = Path(folder) / "cookies.txt"
            valid.write_text(
                "# Netscape HTTP Cookie File\n.example.com\tTRUE\t/\tTRUE\t0\tsession\tvalue\n",
                encoding="utf-8",
            )
            invalid = Path(folder) / "cookies.json"
            invalid.write_text('{"session": "value"}', encoding="utf-8")
            self.assertTrue(validate_cookies_txt(str(valid))[0])
            self.assertFalse(validate_cookies_txt(str(invalid))[0])


class DuplicateTests(unittest.TestCase):
    def test_equivalent_prefixes_are_duplicates(self):
        jobs = [
            gui.parse_line("https://example.com/a"),
            gui.parse_line("gallery-dl https://example.com/a"),
            gui.parse_line("python -m gallery_dl https://example.com/a"),
        ]
        self.assertEqual(gui.find_exact_duplicate_groups(jobs), [[0, 1, 2]])

    def test_different_options_are_not_duplicates(self):
        jobs = [
            gui.parse_line("gallery-dl --range 1-2 https://example.com/a"),
            gui.parse_line("gallery-dl --range 3-4 https://example.com/a"),
        ]
        self.assertEqual(gui.find_exact_duplicate_groups(jobs), [])

    def test_same_command_in_different_tags_is_not_removed(self):
        jobs = gui.parse_text_database(
            "# Project A\n"
            "https://example.com/a\n"
            "# Project B\n"
            "https://example.com/a\n"
        )
        self.assertEqual(gui.find_exact_duplicate_groups(jobs), [])


class WorkerSafetyTests(unittest.TestCase):
    def test_worker_sanitizes_malformed_service_policy_values(self):
        worker = make_worker(
            service_policy={"example": {"delay": "broken", "retries": ["broken"]}},
        )
        job = gui.DownloadJob(
            raw="https://example.com/a",
            is_command=False,
            url="https://example.com/a",
            service="example",
            ident="-",
            dest="-",
        )

        worker.apply_service_delay(job)
        command = worker.build_command(job)

        self.assertNotIn("--retries", command)

    def test_worker_uses_home_relative_config_file(self):
        with tempfile.TemporaryDirectory() as home:
            config = Path(home) / "gallery-dl" / "config.json"
            config.parent.mkdir()
            config.write_text("{}", encoding="utf-8")
            with patch.dict(
                os.environ,
                {"HOME": home, "USERPROFILE": home},
            ):
                command = make_worker(
                    config_path="~/gallery-dl/config.json",
                    output_dir="",
                ).build_command(gui.parse_line("https://example.com/a"))

        config_index = command.index("--config")
        self.assertEqual(command[config_index + 1], str(config))

    def test_probe_timeout_redacts_credentials_from_error_signal(self):
        worker = gui.CommandProbeWorker(
            "gallery-dl",
            ["--password", "probe-secret"],
            timeout=1,
        )
        completed = []
        worker.done.connect(lambda *values: completed.append(values))
        process = MagicMock()
        process.pid = 123
        process.poll.return_value = None
        process.communicate.side_effect = [
            subprocess.TimeoutExpired(
                ["gallery-dl", "--password", "probe-secret"],
                1,
            ),
            ("", ""),
        ]

        with (
            patch("gallery_dl_app.workers.subprocess.Popen", return_value=process),
            patch("gallery_dl_app.workers.IS_WINDOWS", True),
            patch("gallery_dl_app.workers.os.kill"),
        ):
            worker.run()

        self.assertEqual(len(completed), 1)
        self.assertNotIn("probe-secret", completed[0][2])
        self.assertIn(gui.REDACTED, completed[0][2])

    def test_probe_redaction_does_not_replace_short_secret_substrings(self):
        worker = gui.CommandProbeWorker(
            "gallery-dl",
            ["--password", "a"],
            timeout=1,
        )
        completed = []
        worker.done.connect(lambda *values: completed.append(values))
        process = MagicMock()
        process.pid = 123
        process.poll.return_value = None
        process.communicate.side_effect = [
            subprocess.TimeoutExpired(["gallery-dl", "--password", "a"], 1),
            ("", ""),
        ]

        with (
            patch("gallery_dl_app.workers.subprocess.Popen", return_value=process),
            patch("gallery_dl_app.workers.IS_WINDOWS", True),
            patch("gallery_dl_app.workers.os.kill"),
        ):
            worker.run()

        self.assertEqual(len(completed), 1)
        self.assertIn("gallery-dl", completed[0][2])
        self.assertIn(gui.REDACTED, completed[0][2])

    def test_probe_preserves_both_stdout_and_stderr_diagnostics(self):
        worker = gui.CommandProbeWorker("gallery-dl", ["--version"])
        completed = []
        worker.done.connect(lambda *values: completed.append(values))
        process = MagicMock()
        process.communicate.return_value = (
            "startup warning\n",
            "fatal dependency error\n",
        )
        process.returncode = 2

        with patch("gallery_dl_app.workers.subprocess.Popen", return_value=process):
            worker.run()

        self.assertEqual(len(completed), 1)
        self.assertIn("startup warning", completed[0][0])
        self.assertIn("fatal dependency error", completed[0][0])
        self.assertEqual(completed[0][1], 2)

    def test_finished_run_discards_unclaimed_stopped_queue_items(self):
        task_queue = queue.Queue()
        task_queue.put((0, gui.parse_line("https://example.com/a")))
        task_queue.put((1, gui.parse_line("https://example.com/b")))
        harness = MagicMock()
        harness.task_queue = task_queue

        QueueControllerMixin._discard_pending_queue_items(harness)

        self.assertTrue(task_queue.empty())
        self.assertEqual(task_queue.unfinished_tasks, 0)

    def test_rerun_clears_stale_done_state_before_cancel_can_observe_it(self):
        harness = MagicMock()
        harness.results = {0: object(), 1: object()}
        harness.done_indices = {0, 1}
        harness.failed_indices = {2}
        harness.stopped_indices = {1}
        harness.cancelled_indices = {1}

        QueueControllerMixin._prepare_result_sets_for_run(harness, [1], False)

        self.assertEqual(harness.done_indices, {0})
        self.assertEqual(harness.failed_indices, {2})
        self.assertEqual(harness.stopped_indices, set())
        self.assertEqual(harness.cancelled_indices, set())
        self.assertIn(1, harness.results)

    def test_retry_run_status_ignores_failures_from_other_old_rows(self):
        harness = MagicMock()
        harness._processed_in_run = {1}
        harness.failed_indices = {2}
        harness.stopped_indices = set()
        harness.cancelled_indices = set()

        self.assertFalse(QueueControllerMixin._current_run_has_errors(harness))

        harness.failed_indices.add(1)
        self.assertTrue(QueueControllerMixin._current_run_has_errors(harness))

    def test_persistence_failure_is_logged_once_per_operation(self):
        harness = MagicMock()
        harness._persistence_failures_reported = set()

        QueueControllerMixin._report_persistence_failure(
            harness,
            "saving item result",
            RuntimeError("database is locked"),
        )
        QueueControllerMixin._report_persistence_failure(
            harness,
            "saving item result",
            RuntimeError("database is still locked"),
        )

        harness.append_log.assert_called_once()
        self.assertIn("persistence failed", harness.append_log.call_args.args[0])
        self.assertIn("database is locked", harness.append_log.call_args.args[0])

    def test_nonzero_exit_always_counts_at_least_one_error(self):
        command = subprocess.list2cmdline(
            [sys.executable, "-c", "print('plain failure'); raise SystemExit(7)"]
        )
        task_queue = queue.Queue()
        task_queue.put((0, gui.parse_line("https://example.com/a")))
        worker = make_worker(
            task_queue=task_queue,
            gdl_cmd=command,
            output_dir="",
            retries=0,
        )
        completed = []
        worker.job_done.connect(lambda *values: completed.append(values))

        worker.run()

        self.assertEqual(completed[0][2], "failed")
        self.assertEqual(completed[0][5], 1)
        self.assertEqual(completed[0][7], 7)

    def test_config_directory_is_not_passed_as_config_file(self):
        with tempfile.TemporaryDirectory() as folder:
            worker = make_worker(config_path=folder, output_dir="")
            job = gui.parse_line("https://example.com/a")

            command = worker.build_command(job)

            self.assertNotIn("--config", command)

    def test_worker_executes_subprocess_and_finishes_queue_item(self):
        command = subprocess.list2cmdline(
            [sys.executable, "-c", "print('C:/Downloads/file.jpg')"]
        )
        task_queue = queue.Queue()
        task_queue.put((0, gui.parse_line("https://example.com/a")))
        worker = make_worker(
            task_queue=task_queue,
            gdl_cmd=command,
            output_dir="",
            retries=0,
        )
        completed = []
        worker.job_done.connect(lambda *values: completed.append(values))

        worker.run()

        self.assertEqual(task_queue.unfinished_tasks, 0)
        self.assertEqual(len(completed), 1)
        self.assertEqual(completed[0][2], "done")
        self.assertEqual(completed[0][3], 1)

    def test_worker_bounds_oversized_subprocess_line(self):
        command = subprocess.list2cmdline(
            [
                sys.executable,
                "-c",
                f"print('X' * ({gui.MAX_LOG_LINE_CHARS} * 8))",
            ]
        )
        task_queue = queue.Queue()
        task_queue.put((0, gui.parse_line("https://example.com/a")))
        worker = make_worker(
            task_queue=task_queue,
            gdl_cmd=command,
            output_dir="",
            retries=0,
        )
        completed = []
        logs = []
        worker.job_done.connect(lambda *values: completed.append(values))
        worker.log.connect(lambda _worker_id, line: logs.append(line))

        worker.run()

        oversized = [line for line in logs if "[truncated]" in line]
        self.assertEqual(len(oversized), 1)
        self.assertEqual(
            len(oversized[0]),
            gui.MAX_LOG_LINE_CHARS + len(" ... [truncated]"),
        )
        self.assertEqual(completed[0][2], "done")
        self.assertEqual(task_queue.unfinished_tasks, 0)

    def test_cancel_interrupts_service_delay(self):
        cancelled = {7}
        worker = make_worker(
            cancelled_indices=cancelled,
            service_policy={"example": {"delay": 5}},
        )
        job = gui.DownloadJob(
            raw="https://example.com/a",
            is_command=False,
            url="https://example.com/a",
            service="example",
            ident="-",
            dest="-",
        )
        started = time.monotonic()
        worker.apply_service_delay(job, 7)
        self.assertLess(time.monotonic() - started, 0.2)

    def test_worker_status_and_log_hide_url_query_token(self):
        command = subprocess.list2cmdline([sys.executable, "-c", "pass"])
        task_queue = queue.Queue()
        task_queue.put(
            (
                0,
                gui.parse_line("https://example.com/file?access_token=url-secret&item=42"),
            )
        )
        worker = make_worker(
            task_queue=task_queue,
            gdl_cmd=command,
            output_dir="",
            retries=0,
        )
        statuses = []
        logs = []
        worker.status.connect(lambda *_values: statuses.append(_values))
        worker.log.connect(lambda *_values: logs.append(_values))

        worker.run()

        display_text = repr(statuses) + repr(logs)
        self.assertNotIn("url-secret", display_text)
        self.assertIn(gui.REDACTED, display_text)
        self.assertTrue(
            all(ord(character) < 128 for _worker_id, line in logs for character in line),
            logs,
        )

    def test_worker_redacts_credentials_echoed_by_subprocess(self):
        leaked = "https://user:pass@example.com/file?access_token=secret-value"
        command = subprocess.list2cmdline([sys.executable, "-c", f"print({leaked!r})"])
        task_queue = queue.Queue()
        task_queue.put((0, gui.parse_line("https://example.com/a")))
        worker = make_worker(task_queue=task_queue, gdl_cmd=command, output_dir="", retries=0)
        logs = []
        worker.log.connect(lambda *_values: logs.append(_values))

        worker.run()

        display_text = repr(logs)
        self.assertNotIn("user:pass", display_text)
        self.assertNotIn("secret-value", display_text)
        self.assertIn(gui.REDACTED, display_text)

    def test_failed_zip_postprocess_preserves_existing_archive(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "creator"
            destination.mkdir()
            (destination / "image.jpg").write_bytes(b"image")
            archive = Path(folder) / "creator.zip"
            archive.write_bytes(b"existing archive")
            worker = make_worker(compress_enabled=True)
            job = gui.parse_line(f'gallery-dl -d "{destination}" https://example.com/a')

            with patch("gallery_dl_app.workers.zipfile.ZipFile", side_effect=OSError("disk full")):
                result = worker.maybe_compress(job)

            self.assertIn("compression failed", result)
            self.assertEqual(archive.read_bytes(), b"existing archive")
            self.assertEqual(list(Path(folder).glob(".creator.zip.*.tmp")), [])

    def test_zip_postprocess_uses_distinct_temporary_files_per_call(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "creator"
            destination.mkdir()
            (destination / "image.jpg").write_bytes(b"image")
            worker = make_worker(compress_enabled=True)
            job = gui.parse_line(
                f'gallery-dl -d "{destination}" https://example.com/a'
            )
            temporary_paths = []
            with patch(
                "gallery_dl_app.workers.os.replace",
                side_effect=lambda source, _target: temporary_paths.append(Path(source)),
            ):
                worker.maybe_compress(job)
                worker.maybe_compress(job)

        self.assertEqual(len(temporary_paths), 2)
        self.assertNotEqual(temporary_paths[0], temporary_paths[1])

    def test_7z_postprocess_uses_distinct_temporary_files_per_call(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "creator"
            destination.mkdir()
            (destination / "image.jpg").write_bytes(b"image")
            worker = make_worker(compress_enabled=True, compress_format="7z")
            job = gui.parse_line(
                f'gallery-dl -d "{destination}" https://example.com/a'
            )
            process = MagicMock()
            process.communicate.return_value = ("", "")
            process.returncode = 0
            with (
                patch("gallery_dl_app.workers.shutil.which", return_value="7z"),
                patch("gallery_dl_app.workers.subprocess.Popen", return_value=process) as popen,
                patch("gallery_dl_app.workers.os.replace"),
            ):
                worker.maybe_compress(job)
                worker.maybe_compress(job)

        temporary_paths = [Path(call.args[0][3]) for call in popen.call_args_list]
        self.assertEqual(len(temporary_paths), 2)
        self.assertNotEqual(temporary_paths[0], temporary_paths[1])

    def test_postprocess_failure_marks_completed_download_as_failed(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "creator"
            destination.mkdir()
            (destination / "image.jpg").write_bytes(b"image")
            command = subprocess.list2cmdline(
                [sys.executable, "-c", "print('download completed')"]
            )
            task_queue = queue.Queue()
            task_queue.put((0, gui.parse_line(
                f'gallery-dl -d "{destination}" https://example.com/a'
            )))
            worker = make_worker(
                task_queue=task_queue,
                gdl_cmd=command,
                output_dir="",
                retries=0,
                compress_enabled=True,
            )
            completed = []
            worker.job_done.connect(lambda *values: completed.append(values))

            with patch(
                "gallery_dl_app.workers.zipfile.ZipFile",
                side_effect=OSError("disk full"),
            ):
                worker.run()

            self.assertEqual(completed[0][2], "failed")
            self.assertEqual(completed[0][5], 1)
            self.assertEqual(completed[0][7], -1)
            self.assertIn("compression failed", completed[0][8])

    def test_failed_webp_replace_leaves_no_partial_destination(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "creator"
            destination.mkdir()
            png = destination / "image.png"
            Image.new("RGB", (2, 2), "red").save(png)
            worker = make_worker(convert_png_webp=True)
            job = gui.parse_line(f'gallery-dl -d "{destination}" https://example.com/a')
            logs = []
            worker.log.connect(lambda *_values: logs.append(_values))

            with patch("gallery_dl_app.workers.os.replace", side_effect=OSError("disk full")):
                result = worker.maybe_compress(job)

            self.assertFalse((destination / "image.webp").exists())
            self.assertEqual(list(destination.glob(".image.webp.*.tmp")), [])
            self.assertIn("failed", repr(logs).lower())
            self.assertIn("conversion failed", result)

    def test_shared_output_is_not_post_processed_per_worker(self):
        worker = make_worker(compress_enabled=True, convert_png_webp=True)
        job = gui.parse_line("https://example.com/a")
        result = worker.maybe_compress(job)
        self.assertIn("explicit per-job destination", result)


class SpreadsheetImportTests(unittest.TestCase):
    def test_xlsx_rejects_hostile_declared_worksheet_dimensions(self):
        for max_row, max_column, label in (
            (gui.MAX_QUEUE_SOURCE_ROWS + 1, 1, "rows"),
            (1, 16_384, "columns"),
            (10_000, 256, "cells"),
        ):
            with self.subTest(label=label):
                workbook = MagicMock()
                workbook.active.max_row = max_row
                workbook.active.max_column = max_column
                with (
                    patch("gallery_dl_app.queue_controller.validate_xlsx_archive"),
                    patch("openpyxl.load_workbook", return_value=workbook),
                    self.assertRaisesRegex(ValueError, "worksheet|dimension|safety"),
                ):
                    QueueControllerMixin().read_xlsx_as_commands("hostile.xlsx")

                workbook.active.iter_rows.assert_not_called()
                workbook.close.assert_called_once()

    def test_headerless_xlsx_flattens_embedded_newlines_in_command_cell(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "multiline.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["https://example.com/a\n--simulate"])
            workbook.save(path)
            workbook.close()

            text = QueueControllerMixin().read_xlsx_as_commands(str(path))

        jobs = gui.parse_text_database(text)
        self.assertEqual(len(jobs), 1)
        self.assertEqual(
            gui.split_command(jobs[0].raw),
            ["https://example.com/a", "--simulate"],
        )
    def test_csv_limits_disabled_source_rows_during_conversion(self):
        source = "url,enabled\nhttps://example.com/a,false\nhttps://example.com/b,false\n"
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "disabled.csv"
            path.write_text(source, encoding="utf-8")
            with patch("gallery_dl_app.queue_controller.MAX_QUEUE_SOURCE_ROWS", 2):
                with self.assertRaisesRegex(ValueError, "source-row"):
                    QueueControllerMixin().read_csv_as_commands(str(path))

    def test_xlsx_limits_disabled_source_rows_during_conversion(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "disabled.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["url", "enabled"])
            sheet.append(["https://example.com/a", False])
            sheet.append(["https://example.com/b", False])
            workbook.save(path)
            workbook.close()
            with patch("gallery_dl_app.queue_controller.MAX_QUEUE_SOURCE_ROWS", 2):
                with self.assertRaisesRegex(ValueError, "source-row"):
                    QueueControllerMixin().read_xlsx_as_commands(str(path))

    def test_xlsx_expansion_limit_is_checked_before_openpyxl(self):
        fake_member = type("ZipMember", (), {"file_size": MAX_XLSX_UNCOMPRESSED_BYTES + 1})()
        fake_archive = MagicMock()
        fake_archive.__enter__.return_value.infolist.return_value = [fake_member]
        with patch("gallery_dl_app.queue_controller.zipfile.ZipFile", return_value=fake_archive):
            with self.assertRaisesRegex(ValueError, "safety limit"):
                validate_xlsx_archive("oversized.xlsx")

    def test_small_csv_known_header_is_not_imported_as_a_job(self):
        content = (
            "url,enabled\n"
            "https://example.com/disabled,false\n"
            "https://example.com/enabled,true\n"
        )
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "small.csv"
            path.write_text(content, encoding="utf-8")
            text = QueueControllerMixin().read_csv_as_commands(str(path))

        jobs = gui.parse_text_database(text)
        self.assertEqual([job.url for job in jobs], ["https://example.com/enabled"])

    def test_csv_with_only_disabled_rows_stays_empty(self):
        content = "url,enabled\nhttps://example.com/disabled,false\n"
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "disabled.csv"
            path.write_text(content, encoding="utf-8")
            text = QueueControllerMixin().read_csv_as_commands(str(path))

        self.assertEqual(text, "")
        self.assertEqual(gui.parse_text_database(text), [])

    def test_xlsx_false_enabled_cell_skips_row(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "jobs.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["url", "enabled"])
            sheet.append(["https://example.com/disabled", False])
            sheet.append(["https://example.com/enabled", True])
            workbook.save(path)
            workbook.close()

            text = QueueControllerMixin().read_xlsx_as_commands(str(path))

        self.assertNotIn("disabled", text)
        self.assertIn("https://example.com/enabled", text)

    def test_csv_tag_and_notes_survive_import(self):
        content = (
            "url,enabled,tag,notes\n"
            "https://example.com/a,true,Group A,first note\n"
            "https://example.com/b,true,Group B,second note\n"
            "https://example.com/c,true,,untagged note\n"
        )
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "jobs.csv"
            path.write_text(content, encoding="utf-8")
            text = QueueControllerMixin().read_csv_as_commands(str(path))

        jobs = gui.parse_text_database(text)
        self.assertEqual(
            [(job.tag, job.notes) for job in jobs],
            [
                ("Group A", "first note"),
                ("Group B", "second note"),
                ("", "untagged note"),
            ],
        )

    def test_xlsx_tag_and_notes_survive_import(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "jobs.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["url", "enabled", "tag", "notes"])
            sheet.append(["https://example.com/a", True, "Group A", "first note"])
            sheet.append(["https://example.com/b", True, "Group A", "second note"])
            workbook.save(path)
            workbook.close()
            text = QueueControllerMixin().read_xlsx_as_commands(str(path))

        jobs = gui.parse_text_database(text)
        self.assertEqual(
            [(job.tag, job.notes) for job in jobs],
            [("Group A", "first note"), ("Group A", "second note")],
        )


class SystemPathValidationTests(unittest.TestCase):
    def test_command_preview_redacts_secrets_from_builder_exceptions(self):
        harness = MagicMock()
        harness._rebuild_from_text.return_value = True
        harness.jobs = [gui.parse_line("https://example.com/a")]
        harness.final_command_for_job.side_effect = RuntimeError(
            "Command ['gallery-dl', '--password', 'dialog-secret'] failed"
        )

        SystemToolsMixin.command_preview(harness)

        message = harness.show_scroll_message.call_args.args[1]
        self.assertNotIn("dialog-secret", message)
        self.assertIn(gui.REDACTED, message)

    def test_open_selected_destination_expands_home_directory(self):
        with tempfile.TemporaryDirectory() as home, tempfile.TemporaryDirectory() as cwd:
            expected = Path(home) / "chosen-output"
            expected.mkdir()
            previous_cwd = os.getcwd()
            try:
                os.chdir(cwd)
                harness = MagicMock()
                harness.selected_indices.return_value = [0]
                harness.jobs = [MagicMock(dest="~/chosen-output")]
                with (
                    patch.dict(os.environ, {"HOME": home, "USERPROFILE": home}),
                    patch("gallery_dl_app.queue_controller.open_path", return_value=True) as opener,
                ):
                    QueueControllerMixin.open_selected_destination(harness)
            finally:
                os.chdir(previous_cwd)

        opener.assert_called_once_with(expected)
        harness.show_compact_message.assert_not_called()

    def test_config_detection_includes_gallery_dl_windows_home_directory(self):
        with tempfile.TemporaryDirectory() as folder:
            home = Path(folder) / "home"
            appdata = Path(folder) / "appdata"
            expected = home / "gallery-dl" / "config.json"
            expected.parent.mkdir(parents=True)
            expected.write_text("{}", encoding="utf-8")

            with (
                patch.dict(
                    os.environ,
                    {"HOME": str(home), "USERPROFILE": str(home), "APPDATA": str(appdata)},
                ),
                patch("gallery_dl_app.core.IS_WINDOWS", True),
            ):
                detected = gui.detect_config_path()

        self.assertEqual(detected, str(expected))

    def test_open_output_folder_expands_home_like_other_path_actions(self):
        with tempfile.TemporaryDirectory() as home, tempfile.TemporaryDirectory() as cwd:
            previous_cwd = os.getcwd()
            try:
                os.chdir(cwd)
                harness = MagicMock()
                harness.edit_output.text.return_value = "~/chosen-output"
                with (
                    patch.dict(os.environ, {"HOME": home, "USERPROFILE": home}),
                    patch("gallery_dl_app.system_tools.open_path", return_value=True) as opener,
                ):
                    SystemToolsMixin.open_output_folder(harness)
                    expected = Path(home) / "chosen-output"
                    created = expected.is_dir()
            finally:
                os.chdir(previous_cwd)

        opener.assert_called_once_with(expected)
        self.assertTrue(created)

    def test_open_config_rejects_directory_without_opening_it(self):
        with tempfile.TemporaryDirectory() as folder:
            harness = MagicMock()
            harness.config_path = folder
            harness._ui_is_indonesian.return_value = False

            with patch("gallery_dl_app.system_tools.open_path") as opener:
                SystemToolsMixin.open_or_create_config(harness)

        opener.assert_not_called()
        self.assertIn("not a file", harness.show_compact_message.call_args.args[1])

    def test_validate_config_rejects_directory_explicitly(self):
        with tempfile.TemporaryDirectory() as folder:
            harness = MagicMock()
            harness.config_path = folder

            SystemToolsMixin.validate_config(harness)

        self.assertIn("not a file", harness.show_compact_message.call_args.args[1])


class GuideTests(unittest.TestCase):
    def test_command_guide_uses_options_supported_by_gallery_dl_129(self):
        guide = AdvancedToolsMixin.gallery_dl_command_guide_text(
            MagicMock(),
            "English",
        )

        for unsupported in (
            "--date-after",
            "--date-before",
            "--post-range",
            "--child-range",
            "--tags-whitelist",
            "--tags-blacklist",
        ):
            self.assertNotIn(unsupported, guide)
        self.assertIn("-o date-min=2026-01-01", guide)
        self.assertIn("--range 1-20", guide)


class ReportExportTests(unittest.TestCase):
    def test_background_file_task_redacts_exception_before_emitting_signal(self):
        emitted: list[str] = []

        def fail(_stop_event):
            raise RuntimeError("--password background-task-secret")

        task = CancellableFileTask(fail)
        task.failed.connect(emitted.append)

        task.run()

        self.assertEqual(len(emitted), 1)
        self.assertNotIn("background-task-secret", emitted[0])
        self.assertIn(gui.REDACTED, emitted[0])

    def test_compact_dialog_redacts_sensitive_message_text(self):
        harness = MagicMock()
        harness._ui_translate_text.side_effect = lambda value: value
        secret_message = "Command --password dialog-boundary-secret failed"
        with (
            patch("gallery_dl_app.reports.QDialog"),
            patch("gallery_dl_app.reports.QVBoxLayout"),
            patch("gallery_dl_app.reports.QHBoxLayout"),
            patch("gallery_dl_app.reports.QLabel") as label_class,
            patch("gallery_dl_app.reports.QPushButton"),
        ):
            ReportsMixin.show_compact_message(
                harness,
                "Failure",
                secret_message,
                "error",
            )

        rendered = " ".join(
            str(call.args[0]) for call in label_class.call_args_list if call.args
        )
        self.assertNotIn("dialog-boundary-secret", rendered)
        self.assertIn(gui.REDACTED, rendered)

    def test_scroll_dialog_redacts_sensitive_message_text(self):
        harness = MagicMock()
        harness._ui_translate_text.side_effect = lambda value: value
        secret_message = (
            "Command --password scroll-boundary-secret failed\n" + "x" * 800
        )
        with (
            patch("gallery_dl_app.reports.QDialog"),
            patch("gallery_dl_app.reports.QVBoxLayout"),
            patch("gallery_dl_app.reports.QHBoxLayout"),
            patch("gallery_dl_app.reports.QLabel"),
            patch("gallery_dl_app.reports.QPlainTextEdit") as editor_class,
            patch("gallery_dl_app.reports.QPushButton"),
        ):
            ReportsMixin.show_scroll_message(
                harness,
                "Failure",
                secret_message,
                "error",
            )

        rendered = editor_class.return_value.setPlainText.call_args.args[0]
        self.assertNotIn("scroll-boundary-secret", rendered)
        self.assertIn(gui.REDACTED, rendered)

    def test_preflight_rejects_job_destination_that_is_an_existing_file(self):
        with tempfile.TemporaryDirectory() as folder:
            destination = Path(folder) / "not-a-folder"
            destination.write_text("occupied", encoding="utf-8")
            harness = MagicMock()
            harness.edit_output.text.return_value = folder
            harness.jobs = [gui.DownloadJob("url", False, "https://example.com", dest=str(destination))]
            harness.gdl_cmd = "gallery-dl"
            harness.spin_workers.value.return_value = 3

            with patch(
                "gallery_dl_app.reports.QMessageBox.question",
                return_value=QMessageBox.No,
            ) as question:
                allowed = ReportsMixin.run_preflight_audit(harness, [0])

        self.assertFalse(allowed)
        self.assertIn("not a directory", question.call_args.args[2])

    def test_autosave_directory_is_not_moved_as_a_corrupt_file(self):
        with tempfile.TemporaryDirectory() as folder:
            autosave_path = Path(folder) / "autosave_session.json"
            autosave_path.mkdir()
            marker = autosave_path / "keep.txt"
            marker.write_text("keep", encoding="utf-8")
            harness = MagicMock()

            with patch("gallery_dl_app.reports.AUTOSAVE_FILE", autosave_path):
                ReportsMixin._load_autosave_silently(harness)

            self.assertTrue(marker.is_file())
            harness.apply_session_data.assert_not_called()
            self.assertIn("not a file", harness.append_log.call_args.args[0])

    def test_diagnostic_data_does_not_create_missing_output_directory(self):
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "missing" / "downloads"
            harness = MagicMock()
            harness.edit_output.text.return_value = str(output)
            harness.config_path = None
            harness.gdl_cmd = "gallery-dl"
            harness.jobs = []
            harness.done_indices = set()
            harness.failed_indices = set()
            harness.stopped_indices = set()
            harness.cancelled_indices = set()
            harness.spin_workers.value.return_value = 3
            harness.combo_cookies.currentText.return_value = "none"
            harness.spin_retries.value.return_value = 0
            harness.audit_mode = False
            harness.compact_mode = False
            harness.analyze_log_lines.return_value = {}

            with patch("gallery_dl_app.reports.dependency_status", return_value={}):
                data = ReportsMixin.diagnostic_data(harness)

            self.assertFalse(output.exists())
            self.assertTrue(data["output"]["ok"])
            self.assertFalse(data["output"]["exists"])
            self.assertEqual(Path(data["output"]["capacity_checked_at"]), Path(folder))

    def test_diagnostic_data_redacts_credentials_in_gallery_dl_command(self):
        harness = MagicMock()
        harness.edit_output.text.return_value = "."
        harness.config_path = None
        harness.gdl_cmd = "python -m gallery_dl --password audit-secret"
        harness.jobs = []
        harness.done_indices = set()
        harness.failed_indices = set()
        harness.stopped_indices = set()
        harness.cancelled_indices = set()
        harness.spin_workers.value.return_value = 3
        harness.combo_cookies.currentText.return_value = "none"
        harness.spin_retries.value.return_value = 0
        harness.audit_mode = False
        harness.compact_mode = False
        harness.analyze_log_lines.return_value = {}

        with patch("gallery_dl_app.reports.dependency_status", return_value={}):
            data = ReportsMixin.diagnostic_data(harness)

        self.assertNotIn("audit-secret", data["gallery_dl_command"])
        self.assertIn(gui.REDACTED, data["gallery_dl_command"])

    def test_text_exports_redact_inline_credentials(self):
        source = (
            "gallery-dl --password export-secret "
            "https://example.com/a?token=url-secret"
        )
        harness = MagicMock()
        harness.jobs = [gui.parse_line(source)]
        harness.failed_indices = {0}
        harness.stopped_indices = set()
        harness.cancelled_indices = set()
        harness.txt_commands.toPlainText.return_value = "# Project\n" + source
        harness._safe_write_file.return_value = True

        with patch(
            "gallery_dl_app.reports.QFileDialog.getSaveFileName",
            side_effect=[("failed.txt", ""), ("queue.txt", "")],
        ):
            ReportsMixin.export_failed(harness)
            ReportsMixin.export_txt(harness)

        exported = [call.args[1] for call in harness._safe_write_file.call_args_list]
        self.assertEqual(len(exported), 2)
        for text in exported:
            self.assertNotIn("export-secret", text)
            self.assertNotIn("url-secret", text)
            self.assertIn(gui.REDACTED, text)
            self.assertIn("https://example.com/a", text)
        self.assertTrue(exported[1].startswith("# Project\n"))

    def test_log_export_redacts_again_at_file_boundary(self):
        harness = MagicMock()
        harness.all_log_lines = [
            "gallery-dl --password log-export-secret https://example.com/a"
        ]
        harness._safe_write_file.return_value = True

        with patch(
            "gallery_dl_app.reports.QFileDialog.getSaveFileName",
            return_value=("combined_log.txt", ""),
        ):
            ReportsMixin.export_logs(harness)

        exported = harness._safe_write_file.call_args.args[1]
        self.assertNotIn("log-export-secret", exported)
        self.assertIn(gui.REDACTED, exported)

    def test_queue_clipboard_actions_redact_inline_credentials(self):
        source = (
            "gallery-dl --password clipboard-secret "
            "https://example.com/a?token=url-secret"
        )
        harness = MagicMock()
        harness.jobs = [gui.parse_line(source)]
        harness.selected_indices.return_value = [0]
        clipboard = MagicMock()

        with patch(
            "gallery_dl_app.queue_controller.QApplication.clipboard",
            return_value=clipboard,
        ):
            QueueControllerMixin.copy_selected_command(harness)
            QueueControllerMixin.copy_selected_url(harness)

        copied = [call.args[0] for call in clipboard.setText.call_args_list]
        self.assertEqual(len(copied), 2)
        for text in copied:
            self.assertNotIn("clipboard-secret", text)
            self.assertNotIn("url-secret", text)
            self.assertIn(gui.REDACTED, text)


class BackupTests(unittest.TestCase):
    def test_backups_use_distinct_temporary_files_per_call(self):
        with tempfile.TemporaryDirectory() as folder:
            app_dir = Path(folder) / "app-data"
            backup_dir = app_dir / "backups"
            backup_dir.mkdir(parents=True)
            (app_dir / "session.json").write_text("{}", encoding="utf-8")
            target = backup_dir / "backup.zip"
            temporary_paths = []
            with patch(
                "gallery_dl_app.reports.os.replace",
                side_effect=lambda source, _target: temporary_paths.append(Path(source)),
            ):
                write_app_data_backup(app_dir, backup_dir, target)
                write_app_data_backup(app_dir, backup_dir, target)

        self.assertEqual(len(temporary_paths), 2)
        self.assertNotEqual(temporary_paths[0], temporary_paths[1])

    def test_unique_path_does_not_reuse_same_timestamped_name(self):
        with tempfile.TemporaryDirectory() as folder:
            base = Path(folder) / "config.json.bak_20260101_120000"
            base.write_text("first", encoding="utf-8")
            second = base.with_name(base.name + ".2")
            second.write_text("second", encoding="utf-8")

            candidate = gui.unique_path(base)

        self.assertEqual(candidate.name, base.name + ".3")

    def test_app_data_backup_excludes_previous_backup_archives(self):
        with tempfile.TemporaryDirectory() as folder:
            app_dir = Path(folder) / "app-data"
            backup_dir = app_dir / "backups"
            profiles_dir = app_dir / "profiles"
            backup_dir.mkdir(parents=True)
            profiles_dir.mkdir()
            (app_dir / "autosave_session.json").write_text("{}", encoding="utf-8")
            (profiles_dir / "default.json").write_text("{}", encoding="utf-8")
            old_backup = backup_dir / "app_data_backup_old.zip"
            old_backup.write_bytes(b"old backup")
            target = backup_dir / "app_data_backup_new.zip"

            write_app_data_backup(app_dir, backup_dir, target)
            with zipfile.ZipFile(target) as archive:
                relative = set(archive.namelist())

        self.assertEqual(
            relative,
            {"autosave_session.json", "profiles/default.json"},
        )

    def test_failed_app_data_backup_preserves_existing_target(self):
        with tempfile.TemporaryDirectory() as folder:
            app_dir = Path(folder) / "app-data"
            backup_dir = app_dir / "backups"
            backup_dir.mkdir(parents=True)
            (app_dir / "session.json").write_text("{}", encoding="utf-8")
            target = backup_dir / "backup.zip"
            target.write_bytes(b"existing-good-backup")

            with patch("gallery_dl_app.reports.zipfile.ZipFile", side_effect=OSError("disk full")):
                with self.assertRaises(OSError):
                    write_app_data_backup(app_dir, backup_dir, target)

            self.assertEqual(target.read_bytes(), b"existing-good-backup")
            self.assertEqual(list(backup_dir.glob("*.tmp")), [])

    def test_cancelled_app_data_backup_preserves_existing_target(self):
        with tempfile.TemporaryDirectory() as folder:
            app_dir = Path(folder) / "app-data"
            backup_dir = app_dir / "backups"
            backup_dir.mkdir(parents=True)
            (app_dir / "session.json").write_text("{}", encoding="utf-8")
            target = backup_dir / "backup.zip"
            target.write_bytes(b"existing-good-backup")
            stop = threading.Event()
            stop.set()

            with self.assertRaises(InterruptedError):
                write_app_data_backup(app_dir, backup_dir, target, stop)

            self.assertEqual(target.read_bytes(), b"existing-good-backup")
            self.assertEqual(list(backup_dir.glob("*.tmp")), [])

    def test_output_scan_keeps_only_ten_largest_files(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for index in range(15):
                (root / f"file-{index}.bin").write_bytes(b"x" * index)

            result = scan_output_tree(root)

            self.assertEqual(result["files"], 15)
            self.assertEqual(len(result["largest"]), 10)
            self.assertEqual(result["largest"][0][0], 14)


class DropImportTests(unittest.TestCase):
    def test_drop_labels_only_successfully_loaded_files(self):
        with tempfile.TemporaryDirectory() as folder:
            invalid = Path(folder) / "broken.bin"
            invalid.write_bytes(b"not text\x00binary")
            valid = Path(folder) / "working.txt"
            valid.write_text("https://example.com/a\n", encoding="utf-8")

            invalid_url = MagicMock()
            invalid_url.toLocalFile.return_value = str(invalid)
            valid_url = MagicMock()
            valid_url.toLocalFile.return_value = str(valid)
            event = MagicMock()
            event.mimeData.return_value.urls.return_value = [invalid_url, valid_url]
            harness = MagicMock()
            harness.active_workers = 0

            AdvancedToolsMixin.dropEvent(harness, event)

        self.assertEqual(harness.current_file, str(valid))
        harness.lbl_loaded_file.setText.assert_called_once_with("working.txt")
        self.assertEqual(
            harness.txt_commands.setPlainText.call_args.args[0],
            "https://example.com/a",
        )
        harness.show_compact_message.assert_called_once()
        event.acceptProposedAction.assert_called_once_with()

    def test_dry_run_accepts_input_file_and_oauth_commands(self):
        harness = MagicMock()
        harness._rebuild_from_text.return_value = True
        harness.jobs = [
            gui.parse_line("gallery-dl -i urls.txt"),
            gui.parse_line("gallery-dl oauth:pixiv"),
            gui.parse_line("oauth:pixiv"),
        ]
        harness.gdl_cmd = "gallery-dl"
        harness.config_path = None
        harness.edit_output.text.return_value = ""
        harness.combo_cookies.currentText.return_value = "none"
        harness.spin_retries.value.return_value = 0
        harness.service_policy = {}
        validator = MagicMock()
        validator.build_command.side_effect = (
            ["gallery-dl", "-i", "urls.txt"],
            ["gallery-dl", "oauth:pixiv"],
            ["gallery-dl", "oauth:pixiv"],
        )

        with patch(
            "gallery_dl_app.advanced_tools.DownloadWorker",
            return_value=validator,
        ):
            AdvancedToolsMixin.dry_run_validator(harness)

        message = harness.show_scroll_message.call_args.args[1]
        self.assertIn("Problems: 0", message)


class SessionTests(unittest.TestCase):
    def test_string_boolean_values_are_not_treated_as_true(self):
        false_values = (False, 0, "0", "false", "NO", "off", "disabled", "")
        true_values = (True, 1, "1", "true", "YES", "on", "enabled")
        for value in false_values:
            with self.subTest(value=value):
                self.assertFalse(gui.safe_bool(value, True))
        for value in true_values:
            with self.subTest(value=value):
                self.assertTrue(gui.safe_bool(value, False))
        self.assertTrue(gui.safe_bool("unexpected", True))
        self.assertFalse(gui.safe_bool("unexpected", False))

    def test_empty_queue_autosave_still_restores_preferences(self):
        self.assertTrue(is_restorable_autosave({"schema": 3, "commands": "", "theme": "light"}))
        self.assertTrue(is_restorable_autosave({"commands": "", "workers": 5}))
        self.assertFalse(is_restorable_autosave({}))
        self.assertFalse(is_restorable_autosave([]))

    def test_session_does_not_persist_or_restore_secret_bearing_base_command(self):
        harness = MagicMock()
        harness.txt_commands.toPlainText.return_value = "https://example.com/a"
        harness.gdl_cmd = "python -m gallery_dl --password base-secret"
        harness.config_path = None
        harness.current_theme = "dark"
        harness.help_language = "English"
        harness.audit_mode = False
        harness.compact_mode = False
        harness.notifications_enabled = True
        harness.service_policy = {}
        harness.retry_strategy = {}
        harness.active_account_profile_id = None
        harness.clipboard_inbox_enabled = False
        harness.clipboard_allowed_hosts = ""
        harness.close_to_tray = False

        saved = SystemToolsMixin.session_data(harness)

        self.assertEqual(saved["gdl_cmd"], "")
        self.assertNotIn("base-secret", repr(saved))

        restore = MagicMock()
        restore.gdl_cmd = "detected-gallery-dl"
        restore.config_path = None
        restore.current_theme = "dark"
        restore.help_language = "English"
        restore.audit_mode = False
        restore.notifications_enabled = False
        restore.service_policy = {}
        restore.retry_strategy = {}
        restore.active_account_profile_id = None
        restore.clipboard_inbox_enabled = False
        restore.clipboard_allowed_hosts = ""
        restore.close_to_tray = False
        restore.compact_mode = False
        restore.combo_cookies.findText.return_value = 0
        restore.combo_archive.findText.return_value = 0

        SystemToolsMixin.apply_session_data(restore, {
            "commands": "",
            "gdl_cmd": f"python -m gallery_dl --password {gui.REDACTED}",
        })
        self.assertEqual(restore.gdl_cmd, "detected-gallery-dl")

    def test_session_does_not_restore_directory_as_gallery_dl_executable(self):
        with tempfile.TemporaryDirectory() as folder:
            harness = MagicMock()
            harness.gdl_cmd = "detected-gallery-dl"
            harness.config_path = None
            harness.current_theme = "dark"
            harness.help_language = "English"
            harness.audit_mode = False
            harness.notifications_enabled = False
            harness.service_policy = {}
            harness.retry_strategy = {}
            harness.active_account_profile_id = None
            harness.clipboard_inbox_enabled = False
            harness.clipboard_allowed_hosts = ""
            harness.close_to_tray = False
            harness.compact_mode = False
            harness.combo_cookies.findText.return_value = 0
            harness.combo_archive.findText.return_value = 0

            SystemToolsMixin.apply_session_data(harness, {
                "commands": "",
                "gdl_cmd": folder,
            })

        self.assertEqual(harness.gdl_cmd, "detected-gallery-dl")
        self.assertIn("executable is unavailable", harness.append_log.call_args.args[0])

    def test_session_rejects_non_text_core_fields_before_mutating_state(self):
        for field, value in (
            ("commands", ["https://example.com/a"]),
            ("gdl_cmd", {"command": "gallery-dl"}),
            ("config_path", {"path": "config.json"}),
            ("output_dir", ["downloads"]),
        ):
            with self.subTest(field=field):
                harness = MagicMock()
                harness.gdl_cmd = "detected-gallery-dl"
                harness.config_path = "original-config.json"

                with self.assertRaisesRegex(ValueError, field):
                    SystemToolsMixin.apply_session_data(
                        harness,
                        {"commands": "https://example.com/original", field: value},
                    )

                self.assertEqual(harness.gdl_cmd, "detected-gallery-dl")
                self.assertEqual(harness.config_path, "original-config.json")
                harness.txt_commands.setPlainText.assert_not_called()

    def test_version_check_reports_nonzero_return_code_as_failure(self):
        harness = MagicMock()
        harness.gdl_cmd = sys.executable
        harness.version_worker = None
        callbacks = []
        worker = MagicMock()
        worker.done.connect.side_effect = callbacks.append

        with patch("gallery_dl_app.system_tools.CommandProbeWorker", return_value=worker):
            SystemToolsMixin.check_version(harness)
            callbacks[0]("unsupported option", 64, "")

        worker.start.assert_called_once_with()
        title, detail, level = harness.show_compact_message.call_args.args
        self.assertEqual(title, "Check gallery-dl version failed")
        self.assertEqual(detail, "unsupported option")
        self.assertEqual(level, "error")

    def test_version_check_rejects_unrelated_successful_executable_output(self):
        self.assertTrue(is_gallery_dl_version_output("1.29.7"))
        self.assertTrue(is_gallery_dl_version_output("gallery-dl 1.30.0-dev"))
        self.assertFalse(is_gallery_dl_version_output("Microsoft Windows [Version 10]"))

        harness = MagicMock()
        harness.gdl_cmd = sys.executable
        harness.version_worker = None
        callbacks = []
        worker = MagicMock()
        worker.done.connect.side_effect = callbacks.append

        with patch("gallery_dl_app.system_tools.CommandProbeWorker", return_value=worker):
            SystemToolsMixin.check_version(harness)
            callbacks[0]("not gallery-dl", 0, "")

        title, detail, level = harness.show_compact_message.call_args.args
        self.assertEqual(title, "Check gallery-dl version failed")
        self.assertIn("Unexpected --version output", detail)
        self.assertEqual(level, "error")


if __name__ == "__main__":
    unittest.main()
