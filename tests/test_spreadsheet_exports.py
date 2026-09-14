import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from gallery_dl_app.core import parse_text_database
from gallery_dl_app.queue_controller import QueueControllerMixin
from gallery_dl_app.ui_shell import UiShellMixin


class SpreadsheetExportHarness(UiShellMixin):
    def __init__(self) -> None:
        self.messages: list[tuple[str, str, str]] = []

    def show_compact_message(self, title: str, detail: str, level: str) -> None:
        self.messages.append((title, detail, level))


class SpreadsheetExportTests(unittest.TestCase):
    def test_xlsx_atomic_save_uses_distinct_temporary_files_per_call(self):
        with tempfile.TemporaryDirectory() as folder:
            target = Path(folder) / "template.xlsx"
            workbook = MagicMock()
            temporary_paths = []
            with patch(
                "gallery_dl_app.ui_shell.os.replace",
                side_effect=lambda source, _target: temporary_paths.append(Path(source)),
            ):
                UiShellMixin._xlsx_save_atomic(workbook, target)
                UiShellMixin._xlsx_save_atomic(workbook, target)

        self.assertEqual(len(temporary_paths), 2)
        self.assertNotEqual(temporary_paths[0], temporary_paths[1])

    def test_csv_template_escapes_formula_triggers_and_round_trips(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "gallery_dl_database_template.csv"
            harness = SpreadsheetExportHarness()
            with patch(
                "gallery_dl_app.ui_shell.QFileDialog.getSaveFileName",
                return_value=(str(path), "CSV (*.csv)"),
            ):
                harness.export_csv_template()

            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                rows = list(csv.reader(handle))
            imported = QueueControllerMixin().read_csv_as_commands(str(path))

        self.assertEqual(
            rows[0],
            ["URL", "Destination", "Extra Args", "Enabled", "Tag", "Notes", "Command"],
        )
        extra_args_index = rows[0].index("Extra Args")
        self.assertEqual(rows[2][extra_args_index], "'--no-check-certificate")
        self.assertIn("--no-check-certificate", imported)
        self.assertNotIn("'--no-check-certificate", imported)
        self.assertEqual(harness.messages[-1][2], "info")

    def test_template_pack_exports_excel_safe_round_trip_files(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            harness = SpreadsheetExportHarness()
            with patch(
                "gallery_dl_app.ui_shell.QFileDialog.getExistingDirectory",
                return_value=folder,
            ):
                harness.export_template_pack()

            base = Path(folder)
            csv_path = base / "gallery_dl_database_template.csv"
            xlsx_path = base / "gallery_dl_database_template.xlsx"
            workbook = openpyxl.load_workbook(xlsx_path, data_only=False)
            try:
                formula_cells = [
                    f"{sheet.title}!{cell.coordinate}"
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.data_type == "f"
                ]
            finally:
                workbook.close()

            csv_jobs = parse_text_database(QueueControllerMixin().read_csv_as_commands(str(csv_path)))
            xlsx_jobs = parse_text_database(QueueControllerMixin().read_xlsx_as_commands(str(xlsx_path)))

        self.assertEqual(formula_cells, [])
        self.assertEqual(len(csv_jobs), 3)
        self.assertEqual(len(xlsx_jobs), 3)
        self.assertIn("--no-check-certificate", csv_jobs[1].raw)
        self.assertIn("--no-check-certificate", xlsx_jobs[1].raw)
        self.assertEqual(harness.messages[-1][2], "info")

    def test_xlsx_database_writer_keeps_equals_prefixed_cells_literal(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "literal.xlsx"
            workbook = openpyxl.Workbook()
            UiShellMixin._xlsx_write_database_sheet(
                workbook.active,
                [["command"], ["=not-a-formula"]],
            )
            workbook.save(path)
            workbook.close()

            reopened = openpyxl.load_workbook(path, data_only=False)
            try:
                cell = reopened.active["A2"]
                self.assertEqual(cell.value, "=not-a-formula")
                self.assertEqual(cell.data_type, "s")
            finally:
                reopened.close()

    def test_xlsx_template_help_is_literal_text_not_formulas(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "gallery_dl_database_template.xlsx"
            harness = SpreadsheetExportHarness()
            with patch(
                "gallery_dl_app.ui_shell.QFileDialog.getSaveFileName",
                return_value=(str(path), "Excel (*.xlsx)"),
            ):
                harness.export_xlsx_template()

            workbook = openpyxl.load_workbook(path, data_only=False)
            try:
                formula_cells = [
                    f"{sheet.title}!{cell.coordinate}"
                    for sheet in workbook.worksheets
                    for row in sheet.iter_rows()
                    for cell in row
                    if cell.data_type == "f"
                ]
            finally:
                workbook.close()

            imported = QueueControllerMixin().read_xlsx_as_commands(str(path))
            jobs = parse_text_database(imported)

        self.assertEqual(formula_cells, [])
        self.assertEqual(len(jobs), 3)
        self.assertIn("--no-check-certificate", jobs[1].raw)
        self.assertEqual(harness.messages[-1][2], "info")


if __name__ == "__main__":
    unittest.main()
